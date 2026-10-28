"""Build the university-format abstract from one text source (no experiments).

Run: PYTHONPATH=.cache/graduation_tools .venv/bin/python docs/midterm2026/build_abstract.py
Then compile paper.tex with LuaLaTeX twice. See README.md for verification.
"""
from pathlib import Path
import csv
import json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
SOURCE = ROOT / 'docs/presentation_sources/20261016T_reachable_rescues'
TITLE = '量子LLM支援プログラミングにおける\n失敗類型と介入カタログの実証研究'
AUTHOR = '2023TC011 伊藤 悠真　　指導教員：横山 哲郎'
BODY_FONT = 'ＭＳ Ｐ明朝'
HEADING_FONT = 'ＭＳ Ｐゴシック'
LATIN_FONT = 'Times New Roman'
CONDITIONS = ['self_refine', 'execution_feedback', 'self_debugging']
MODEL_NAMES = [('G3', 'gemma3:4b'), ('G4', 'gemma4:e4b-it-qat'),
               ('C7', 'qwen2.5-coder:7b'), ('Q4', 'qwen3.5:4b'), ('Q9', 'qwen3.5:9b')]
ERROR_NAMES = [('回路構築', 'build_error'), ('import', 'import_error'),
               ('IF不適合', 'interface_mismatch'), ('量子ビット数', 'qubit_count_mismatch'),
               ('構文', 'syntax'), ('実行時', 'unknown_error'), ('出力不一致', 'wrong_output')]
ROWS = list(csv.DictReader((ROOT / 'docs/analysis/reachable_rescues_20261016/core/repair_summary.csv').open()))
B4_SOURCE = ROOT / 'docs/analysis/independent_4_completed_20261022'
B4_ROWS = list(csv.DictReader((B4_SOURCE / 'summary.csv').open()))

def b4_value(grouping, value):
    return next(r['successes_from_initial_failure'] for r in B4_ROWS if r['condition'] == 'B4' and r['generation_cap'] == '4' and r['grouping'] == grouping and r['group'] == value)

SNAP = json.loads((SOURCE / 'results_snapshot.json').read_text())


def rows(grouping, names):
    result = []
    for label, value in names:
        rr = [next(r for r in ROWS if r['scope'] == 'pooled' and r['grouping'] == grouping
                   and r['condition_id'] == c
                   and r['model_spec' if grouping == 'model' else 'baseline_error_category']
                   == ('ollama:' + value if grouping == 'model' else value)) for c in CONDITIONS]
        assert len({r['initial_failures'] for r in rr}) == 1
        result.append([label, rr[0]['initial_failures']] + [r['rescued'] for r in rr] + [b4_value('model' if grouping == 'model' else 'initial_failure', 'ollama:' + value if grouping == 'model' else value)])
    return result


TABLES = [
    ('モデル別の初期失敗数と追加成功数', ['モデル', '初期失敗数', 'R', 'F', 'D', 'B4'], rows('model', MODEL_NAMES)),
    ('初期失敗類型別の追加成功数', ['失敗類型', '初期失敗数', 'R', 'F', 'D', 'B4'], rows('baseline_error', ERROR_NAMES)),
    ('FとDの救済対象（初期L3未達545組）', ['共通', 'Fのみ', 'Dのみ', '双方未救済'],
     [[str(SNAP['primary']['fd_overlap'][k]) for k in ['both', 'F_only', 'D_only', 'neither']]])
]

# Tables are kept in the reading order; the same prose is used in TeX and Word.
BLOCKS = [
('h', 'はじめに'),
('p', '大規模言語モデル（LLM）の発展に伴いコード生成が進み，HumanEval［1］やMBPP［2］で性能が評価されている．一方，生成コードには構文上・意味上の誤りや，存在しないAPIを用いるハルシネーションが生じ得る．量子プログラミングは量子ビットの重ね合わせ・量子もつれと確率的な測定を扱うため，Pythonとして実行できても期待する量子状態・測定分布を得られるとは限らない．段階的な正しさの検証が必要である．'),
('p', '本研究は，Qiskitコードの初期L3未達に対する修正介入による救済率と救済対象の違いを明らかにし，失敗類型に応じた介入カタログの根拠を整理することを目的とする．'),
('h', '評価の枠組み'),
('p', '現行評価器では，L0は構文・importの検証を含むコードの読み込みと，必須関数の要件への適合，L1は回路構築と評価処理の実行，L2は指定入力での期待出力との一致，L3はL0〜L2と構造要件への適合とする．L3の要件をすべて満たすことを「L3成功」と呼ぶ．関数名・戻り値・出力形式などのインターフェース適合性と，量子的な出力一致を区別する．'),
('p', '構造要件は量子ビット数や，測定・barrierを除く命令数の範囲などである．L3は限定された基準であり，全入力のアルゴリズム同値性を保証しない．測定分布は8,192 shotsの全変動距離，状態ベクトルは大域位相を整合した距離を用い，0.05未満を一致とする．測定結果を，測定先の対応関係に基づいて量子ビット順に並べ替える．初期失敗数は初期L3未達の組数とする．追加成功率は「初期L3未達から最終L3成功となった組数／初期失敗数」とし，失敗類型は介入前で固定する．'),
('h', '関連研究・準備'),
('p', 'Self-Refine［3］は自己フィードバックと反復修正，Self-Debugging［4］はコード説明や実行結果を用いる．推論課題では，外部フィードバックのない自己修正の限界も報告されている［5］．QiskitはPythonで量子回路を構築する基盤であり，Qiskit HumanEval［6］はそのコード生成を評価する．本研究では独自の10タスクに修正条件を実装し，先行研究の完全再現とは位置づけない．'),
('h', '実験方法'),
('p', '5モデル・10タスク・seed 0〜19の1,000初期組を対象に，B，R，F，D，Pの5条件，計5,000件を分析した．さらに，独立再生成B4を追加の対照条件として別実験で評価した．モデル略号はQ4：qwen3.5:4b，G4：gemma4:e4b-it-qat，Q9：qwen3.5:9b，C7：qwen2.5-coder:7b，G3：gemma3:4bである．'),
('p', 'タスクはBell，GHZ，Deutsch–Jozsa（constant/balanced），Bernstein–Vazirani，Grover，QFT，IQFT，QPE，Ansatzとした．temperatureは0.7，1要求の生成上限は1,024 tokensである．'),
('p', 'Bは初期生成，Rは自己見直し，Fはエラー等の実行結果提示，Dは説明・原因分析を促す自己デバッグ，Pは生成前注意事項である．R/F/DはBと同じ初期コード・判定から最大3回修正し，成功時に停止する．Pは別プロンプトによる独立生成である．'),
('p', '検証済みの保存記録と部分追試から再構成した生成・修正経路を集計した［7］．全件を新規生成・再実行した結果ではない．追加のB4は同じ初期コードを保持し，初期L3未達の組だけ課題仕様から最大3回独立再生成する．前回コード・診断を与えず，独立再生成用に定めたseed設定規則と成功時停止を用いる．'),
('h', '結果'),
('p', '初期L3成功は455組，未達は545組であった．救済はRが96/545組（17.6%），Fが183/545組（33.6%），Dが180/545組（33.0%）であった．全体の救済率は初期L3未達545組を分母とする合算値である．'),
('t', 0),
('p', '表1・2の追加成功は，R/F/Dでは修正による救済，B4では独立再生成による成功を指す．R/F/Dの追加成功率を救済率と呼ぶ．G4ではFの成功数がDを上回る一方，Q4ではDがFを上回った．'),
('t', 1),
('p', '表2のimport失敗はF 58/271組（21.4%），D 72/271組（26.6%），回路構築失敗はF 76/126組（60.3%），D 68/126組（54.0%）であった．IFはインターフェースを表す．「実行時」の1組は未定義変数simによるNameErrorで，実行時エラー（未定義変数）に分類した．少数分母の率は参考値とする．'),
('t', 2),
('p', 'FとDの全体の救済数は3組差だが，表3の片側のみ救済した83組では対象が異なる．import失敗はFのみ8組・Dのみ22組，回路構築失敗は19組・11組であった．83組すべての両条件166履歴・398修正遷移を静的に確認した．複数箇所が同時に変わるため，変更内容から救済原因は分離できない．少なくとも一方が救済した223/545組（40.9%）は事後的な集合であり，実行前の選択方策の成功率ではない．'),
('p', 'B4の追加実験［7］は1,000組を完了し，欠落・重複・予定外0，sanity検証成功であった．初期L3未達から217/545組（39.8%）が成功し，最終成功は672/1,000組（67.2%）だった．Fより34組，Dより37組多い．B4とFの共通成功は131組，B4のみ86組，Fのみ52組だった．'),
('p', '追加計算コストとして，共有初期生成を除き未成功例も含む追加tokens／追加成功はB4 4,589.0，R 18,852.6，F 9,333.8，D 10,976.6だった．B4では追加で1,342回の生成要求を行い，995,815 tokensを消費した．R/F/Dは保存記録と部分追試からの再構成値である．金銭費用ではない．'),
('h', '考察'),
('p', 'FとDの救済率が近くても救済対象は一致せず，初期失敗類型・モデルによって成功数の関係が異なった．介入カタログには，全体の救済率だけでなく，対象となる失敗類型，初期失敗数，追加計算コストを併記する必要がある．ただし，失敗類型はモデル・タスクと関連しており，この結果だけで介入選択規則を確定できない．'),
('p', 'B4とR/F/Dは初回込み最大4回だが，成功時停止により実要求数は異なる．入力情報・tokens・seed規則・実行時期・出典schemaに差があり，診断提示の因果効果は断定しない．成功時停止，初期コード・タスク・seedの依存，元実験の診断文差もあり，5,000条件を独立標本として統計的有意差や一般的優劣を主張しない．L3の入力範囲，有限shots，生成長上限も限界である．'),
('h', 'おわりに'),
('p', '初期L3未達545組に対し，Fは33.6%，Dは33.0%を救済したが，片側のみ救済した83組の内訳は異なった．B4は初期L3未達545組のうち39.8%を追加成功に導き，再生成対照を含めた検討の必要性を示した．カタログには修正介入に加えて独立再生成も候補として記録する．介入選択・停止方策を事前固定し，別seed・タスクでも検証する．'),
('ref', 'M. Chen et al.：Evaluating Large Language Models Trained on Code，arXiv:2107.03374（2021）．'),
('ref', 'J. Austin et al.：Program Synthesis with Large Language Models，arXiv:2108.07732（2021）．'),
('ref', 'A. Madaan et al.：Self-Refine: Iterative Refinement with Self-Feedback，arXiv:2303.17651（2023）．'),
('ref', 'X. Chen et al.：Teaching Large Language Models to Self-Debug，arXiv:2304.05128（2023）．'),
('ref', 'J. Huang et al.：Large Language Models Cannot Self-Correct Reasoning Yet，arXiv:2310.01798（2023）．'),
('ref', 'S. Vishwakarma et al.：Qiskit HumanEval: An Evaluation Benchmark For Quantum Code Generative Models，arXiv:2406.14712（2024）．'),
('ref', '現行到達経路（2026-10-16版）とB4追加実験（2026-10-22固定）の集計・保存設定・出典hash．本原稿READMEに出典と検証記録を示す．'),
]


def escape(s):
    return ''.join({'&': r'\&', '%': r'\%', '_': r'\_', '#': r'\#'}.get(c, c) for c in s)


def build_tex():
    pre = r'''% Generated by build_abstract.py; prose is shared with paper.docx.
\documentclass[a4paper,10pt,twocolumn]{extarticle}
\usepackage{luatexja-fontspec}
% Required fonts are preferred. Fallback output is a preview, not final.
\IfFontExistsTF{Times New Roman}{
  \setmainfont{Times New Roman}\setsansfont{Times New Roman}
}{
  \PackageWarningNoLine{abstract-fonts}{Times New Roman unavailable; preview uses Nimbus Roman}
  \setmainfont{Nimbus Roman}\setsansfont{Nimbus Roman}
}
\IfFontExistsTF{MS PMincho}{
  \setmainjfont[Scale=1]{MS PMincho}
}{\IfFontExistsTF{ＭＳ Ｐ明朝}{
  \setmainjfont[Scale=1]{ＭＳ Ｐ明朝}
}{
  \PackageWarningNoLine{abstract-fonts}{MS PMincho unavailable; preview uses IPAPMincho}
  \setmainjfont[Scale=1]{IPAPMincho}
}}
\IfFontExistsTF{MS PGothic}{
  \setsansjfont[Scale=1]{MS PGothic}
}{\IfFontExistsTF{ＭＳ Ｐゴシック}{
  \setsansjfont[Scale=1]{ＭＳ Ｐゴシック}
}{
  \PackageWarningNoLine{abstract-fonts}{MS PGothic unavailable; preview uses IPAPGothic}
  \setsansjfont[Scale=1]{IPAPGothic}
}}
\usepackage[left=18mm,right=18mm,top=18mm,bottom=24mm,columnsep=7mm,footskip=12mm]{geometry}
\usepackage{titlesec}
\usepackage{caption}
\usepackage{array}
\usepackage{flushend}
\renewcommand{\tablename}{表}
\renewcommand{\thetable}{\arabic{table}}
\makeatletter
\renewcommand\normalsize{\@setfontsize\normalsize{11bp}{14.5bp}}
\makeatother
\normalsize
\titleformat{\section}{\fontsize{11bp}{14.5bp}\selectfont\sffamily\gtfamily\bfseries}{\thesection}{.65em}{}
\titlespacing*{\section}{0pt}{5bp}{2bp}
\captionsetup{font=normalsize,skip=3bp,labelsep=quad}
\setlength{\parindent}{11bp}
\setlength{\parskip}{0pt}
\setlength{\tabcolsep}{3bp}
\setlength{\intextsep}{5bp}
\pagestyle{plain}
\raggedbottom
\hyphenpenalty=10000
\exhyphenpenalty=10000
\emergencystretch=1em
\begin{document}
\twocolumn[\begin{center}
{\fontsize{11bp}{14bp}\selectfont\sffamily\gtfamily\bfseries TITLE\par}
\vspace{4bp}
AUTHOR
\end{center}\vspace{4bp}]
'''
    out = [pre.replace('TITLE', escape(TITLE).replace('\n', r'\\')).replace('AUTHOR', AUTHOR)]
    refnum = 0
    for kind, text in BLOCKS:
        if kind == 'h':
            if text == '実験方法':
                out.append(r'\newpage')  # Begin the methods in the right column.
            out.append(r'\section{' + text + '}')
        elif kind == 'p':
            out.append(escape(text) + '\n')
        elif kind == 't':
            caption, headers, rr = TABLES[text]
            out.append(r'\begin{center}\begin{minipage}{\linewidth}\centering')
            out.append(r'\captionof{table}{' + caption + '}')
            out.append(r'\begin{tabular}{@{}l' + 'r' * (len(headers)-1) + r'@{}}\hline')
            out.append(' & '.join(headers) + r'\\\hline')
            for row in rr:
                out.append(' & '.join(map(escape, row)) + r'\\')
            out.append(r'\hline\end{tabular}\end{minipage}\end{center}')
        elif kind == 'ref':
            refnum += 1
            if refnum == 1:
                out.append(r'\section*{参考文献}')
            out.append(r'{\raggedright\noindent\hangindent=1.5em ［' + str(refnum) + '］' + escape(text) + r'\par}' + '\n')
    out.append(r'\end{document}')
    (HERE / 'paper.tex').write_text('\n'.join(out) + '\n')


def font(run, size=11, bold=False):
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + name), (HEADING_FONT if bold else BODY_FONT) if name == 'eastAsia' else LATIN_FONT)
        # Explicit font names must take precedence over the document theme.
        rf.attrib.pop(qn('w:' + name + 'Theme'), None)


def build_docx():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Mm(18)
    s.top_margin, s.bottom_margin = Mm(18), Mm(24)
    s.footer_distance = Mm(12)
    st = d.styles['Normal']
    st.font.name, st.font.size = LATIN_FONT, Pt(11)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), BODY_FONT)
    for name in ['Normal', 'Body Text', 'Caption', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = d.styles[name]
        style.font.name, style.font.size = LATIN_FONT, Pt(11)
        rf = style.element.get_or_add_rPr().get_or_add_rFonts()
        for script in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
            rf.set(qn('w:' + script), (HEADING_FONT if name == 'Title' or name.startswith('Heading ') else BODY_FONT)
                   if script == 'eastAsia' else LATIN_FONT)
            rf.attrib.pop(qn('w:' + script + 'Theme'), None)
    pf = st.paragraph_format
    pf.line_spacing, pf.space_after, pf.space_before = Pt(14.5), Pt(0), Pt(0)
    pf.widow_control = True
    # PAGE remains editable and updates after repagination.
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    r = p.add_run('1')
    font(r)
    field.append(r._r)
    p._p.append(field)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = Pt(14)
    font(p.add_run(TITLE), 11, True)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(AUTHOR))
    s = d.add_section(WD_SECTION_START.CONTINUOUS)
    cols = s._sectPr.find(qn('w:cols'))
    cols.set(qn('w:num'), '2')
    cols.set(qn('w:space'), str(round(7 / 25.4 * 1440)))
    secnum = refnum = 0
    for kind, text in BLOCKS:
        if kind in ['h', 'p', 'ref']:
            p = d.add_paragraph()
            if kind == 'h':
                secnum += 1
                text = f'{secnum}　{text}'
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(2)
            elif kind == 'p':
                p.paragraph_format.first_line_indent = Pt(11)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            else:
                refnum += 1
                if refnum == 1:
                    r = p.add_run('参考文献')
                    font(r, 11, True)
                    p.paragraph_format.keep_with_next = True
                    p.paragraph_format.space_before = Pt(5)
                    p = d.add_paragraph()
                text = f'［{refnum}］{text}'
                p.paragraph_format.left_indent = Pt(15)
                p.paragraph_format.first_line_indent = Pt(-15)
            font(p.add_run(text), 11, kind == 'h')
        elif kind == 't':
            caption, headers, rr = TABLES[text]
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            font(p.add_run(f'表{text+1}　{caption}'))
            t = d.add_table(rows=0, cols=len(headers))
            t.autofit = False
            widths = [27, 20, 8, 8, 8, 12.5] if len(headers) == 6 else [19, 19, 19, 26.5]
            for c, w in zip(t.columns, widths):
                c.width = Mm(w)
            tw = t._tbl.tblPr.find(qn('w:tblW'))
            tw.set(qn('w:type'), 'dxa')
            tw.set(qn('w:w'), str(round(83.5 / 25.4 * 1440)))
            margins = OxmlElement('w:tblCellMar')
            for edge in ['top', 'bottom', 'left', 'right']:
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), '0')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            t._tbl.tblPr.append(margins)
            for i, row in enumerate([headers] + rr):
                tr = t.add_row()
                tr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for j, (c, val, w) in enumerate(zip(tr.cells, row, widths)):
                    c.width = Mm(w)
                    p = c.paragraphs[0]
                    p.paragraph_format.keep_with_next = i != len(rr)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                    font(p.add_run(str(val)))
                    borders = OxmlElement('w:tcBorders')
                    for edge in (['top', 'bottom'] if i == 0 else ['bottom'] if i == len(rr) else []):
                        el = OxmlElement('w:' + edge)
                        el.set(qn('w:val'), 'single')
                        el.set(qn('w:sz'), '4')
                        borders.append(el)
                    c._tc.get_or_add_tcPr().append(borders)
    # End the two-column section continuously so Word balances its final page.
    s = d.add_section(WD_SECTION_START.CONTINUOUS)
    s._sectPr.find(qn('w:cols')).set(qn('w:num'), '1')
    d.core_properties.title = TITLE.replace('\n', '')
    d.core_properties.author = '伊藤 悠真'
    d.save(HERE / 'paper.docx')


if __name__ == '__main__':
    build_tex()
    build_docx()
    print('Generated paper.tex and paper.docx from the same text and frozen CSV tables.')
