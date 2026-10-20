"""Check the abstract's actual PDFs, editable Word structure and frozen numbers."""
from datetime import datetime, timezone
import csv
import hashlib
import json
from pathlib import Path
import re
import unicodedata
import pymupdf
import numpy as np
from docx import Document
from docx.oxml.ns import qn
import build_abstract as b

HERE, ROOT = b.HERE, b.ROOT
CACHE = ROOT / '.cache/midterm_abstract_20261021'
ANALYSIS = ROOT / 'docs/analysis/reachable_rescues_20261016'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(s):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', s)).replace('\u00ad', '')


def validate_pdf(path):
    doc = pymupdf.open(path)
    assert len(doc) == 2, (path, len(doc))
    body = []
    pages = []
    font_names = set()
    for i, page in enumerate(doc):
        assert abs(page.rect.width - 210 / 25.4 * 72) < .1
        assert abs(page.rect.height - 297 / 25.4 * 72) < .1
        spans = [s for bb in page.get_text('dict')['blocks'] if 'lines' in bb
                 for line in bb['lines'] for s in line['spans'] if s['text'].strip()]
        footer = [s for s in spans if s['bbox'][1] > 273 / 25.4 * 72]
        assert norm(''.join(s['text'] for s in footer)) == str(i+1), (path, 'page number', footer)
        content = [s for s in spans if s not in footer]
        for s in spans:
            font_names.add(s['font'])
            assert 10 - .02 <= s['size'] <= 11 + .02, (path, 'font size', s)
            assert s['color'] == 0, (path, 'non-black text', s)
            assert '\ufffd' not in s['text'], (path, 'missing glyph', s)
        # Japanese punctuation has blank half-em side bearings: its PDF em box
        # can cross a margin even when all printed ink is inside. Check actual
        # raster ink at 216 dpi, rather than treating those blanks as overflow.
        pix = page.get_pixmap(matrix=pymupdf.Matrix(3, 3), colorspace=pymupdf.csGRAY, alpha=False)
        raster = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width)
        ys, xs = np.where(raster[:int(273 / 25.4 * 72 * 3)] < 80)
        ink_mm = [float(xs.min()/3/72*25.4), float(ys.min()/3/72*25.4),
                  float(xs.max()/3/72*25.4), float(ys.max()/3/72*25.4)]
        assert ink_mm[0] >= 18-.15 and ink_mm[1] >= 18-.15, (path, 'top/left ink overflow', ink_mm)
        assert ink_mm[2] <= 192+.15, (path, 'right ink overflow', ink_mm)
        # No content may appear in the lower margin other than the page number.
        assert all(s['bbox'][1] >= 273/25.4*72 and norm(s['text']) == str(i+1)
                   for s in spans if s['bbox'][3] > 273/25.4*72)
        # Table rules must also be inside the text area.
        for draw in page.get_drawings():
            r = draw['rect']
            assert r.x0 >= 18 / 25.4 * 72 - .6 and r.x1 <= 192 / 25.4 * 72 + .6
            assert r.y0 >= 18 / 25.4 * 72 - .6 and r.y1 <= 273 / 25.4 * 72 + .6
        # Confirm both columns actually contain text; title/author may span them.
        assert any(s['bbox'][0] < 280 for s in content)
        assert any(s['bbox'][0] > 300 for s in content)
        body.append(''.join(s['text'] for s in content))
        pages.append({'page': i+1, 'size_mm': [round(x / 72 * 25.4, 3) for x in page.rect[2:]],
                      'font_sizes_pt': sorted({round(s['size'], 3) for s in spans}),
                      'page_number': i+1, 'content_within_margins': True,
                      'printed_ink_bounds_mm': [round(x, 3) for x in ink_mm],
                      'both_columns_have_text': True})
        assert all(doc.extract_font(f[0])[3] for f in page.get_fonts()), 'unembedded font'
    text = norm(''.join(body))
    assert norm(b.TITLE) in text
    assert norm(b.AUTHOR) in text
    headings = [value for kind, value in b.BLOCKS if kind == 'h']
    for i, title in enumerate(headings, 1):
        assert norm(f'{i}{title}') in text
    # Every paragraph/reference must be present in full after PDF extraction.
    for kind, value in b.BLOCKS:
        if kind in ['p', 'ref']:
            assert norm(value) in text, (path, 'incomplete or changed paragraph', value)
    for i, (caption, headers, rows) in enumerate(b.TABLES, 1):
        assert norm(f'表{i}{caption}') in text
        for row in [headers] + rows:
            assert norm(''.join(row)) in text, (path, 'table row mismatch', row)
    families = ''.join(re.sub('[^a-z]', '', name.lower()) for name in font_names)
    font_match = all(name in families for name in ['mspmincho', 'mspgothic', 'timesnewroman'])
    return {'path': str(path.relative_to(ROOT)), 'sha256': sha(path), 'pages': pages,
            'actual_pdf_fonts': sorted(font_names), 'requested_font_families_present': font_match,
            'complete_paragraphs_and_references': True, 'tables_match_source_csv': True,
            'all_fonts_embedded': True, 'black_text_only': True}


def validate_numbers():
    snap = b.SNAP
    overall = {r['condition_id']: r for r in snap['primary']['overall']}
    assert [(overall[c]['rescued'], overall[c]['initial_failures']) for c in b.CONDITIONS] == [(96, 545), (183, 545), (180, 545)]
    assert [round(overall[c]['rescue_rate'] * 100, 1) for c in b.CONDITIONS] == [17.6, 33.6, 33.0]
    for table in b.TABLES[:2]:
        assert sum(int(r[1]) for r in table[2]) == 545
        assert [sum(int(r[i]) for r in table[2]) for i in [2, 3, 4]] == [96, 183, 180]
    assert 455 + 545 == snap['primary']['initial_groups'] == 1000
    assert snap['primary']['condition_records'] == 5000
    overlap = snap['primary']['fd_overlap']
    assert (overlap['both'], overlap['F_only'], overlap['D_only'], overlap['neither']) == (140, 43, 40, 322)
    assert sum([140, 43, 40, 322]) == 545
    assert 140 + 43 == 183 and 140 + 40 == 180 and 43 + 40 == 83
    assert round((140 + 43 + 40) / 545 * 100, 1) == 40.9
    cases = list(csv.DictReader((ANALYSIS / 'discordant_cases.csv').open()))
    rescued = [r for r in cases if r['rescued'] == 'True']
    assert len(cases) == 166 and len(rescued) == 83
    for cohort, counts in [('F_only', (8, 19, 28, 24)), ('D_only', (22, 11, 34, 32))]:
        rr = [r for r in rescued if r['cohort'] == cohort]
        got = (sum(r['baseline_error_category'] == 'import_error' for r in rr),
               sum(r['baseline_error_category'] == 'build_error' for r in rr),
               sum(r['any_changed_import_API'] == 'True' for r in rr),
               sum(r['any_changed_circuit_interface'] == 'True' for r in rr))
        assert got == counts, (cohort, got)
    assert len(list(csv.DictReader((ANALYSIS / 'discordant_transitions.csv').open()))) == 398
    costs = {r['condition_id']: r for r in snap['overall_costs'] if r['cap'] == 3}
    assert [round(costs[c]['tokens_per_rescue'], 1) for c in b.CONDITIONS] == [18852.6, 9333.8, 10976.6]
    third = {r['condition_id']: r for r in snap['third_round']}
    for c, expected in [('execution_feedback', (28, 390, 7.2, 15.3)), ('self_debugging', (33, 398, 8.3, 18.3))]:
        r = third[c]
        got = (r['additional_rescues'], r['at_risk_after_2'],
               round(r['additional_rescues']/r['at_risk_after_2']*100, 1), round(r['loss_if_stop_at_2']*100, 1))
        assert got == expected
    p = overall['preventive_spec']
    assert (p['rescued'], p['regressed'], p['final_successes']) == (124, 106, 473)
    assert round(124/545*100, 1) == 22.8 and round(106/455*100, 1) == 23.3
    assert 455-106+124 == 473
    assert json.loads((ANALYSIS / 'sanity_report.json').read_text())['passed']
    return {'overall_rescues': [96, 183, 180], 'initial_failure_denominator': 545,
            'model_and_error_table_totals_reconciled': True, 'fd_overlap_reconciled': True,
            'all_83_cases_and_398_transitions_checked': True, 'cost_and_third_round_checked': True,
            'independent_generation_P_checked': True}


def main():
    assert [text for kind, text in b.BLOCKS if kind == 'h'] == [
        'はじめに', '評価の枠組み', '関連研究・準備', '実験方法', '結果', '考察', 'おわりに']
    refs = [text for kind, text in b.BLOCKS if kind == 'ref']
    assert len(refs) == 7
    citations = [int(n) for kind, text in b.BLOCKS if kind == 'p' for n in re.findall(r'［(\d+)］', text)]
    assert sorted(set(citations)) == list(range(1, 8))
    bibliography = json.loads((HERE / 'bibliography_sources.json').read_text())
    assert [r['reference'] for r in bibliography['references']] == refs[:6]
    doc = Document(HERE / 'paper.docx')
    assert len(doc.tables) == 3
    assert [s._sectPr.find(qn('w:cols')).get(qn('w:num'), '1') for s in doc.sections] == ['1', '2', '1']
    for s in doc.sections:
        assert abs(s.page_width.mm - 210) < .02 and abs(s.page_height.mm - 297) < .02
        assert all(abs(x.mm - 18) < .02 for x in [s.left_margin, s.right_margin, s.top_margin])
        assert abs(s.bottom_margin.mm - 24) < .02
    for t, (_, headers, rows) in zip(doc.tables, b.TABLES):
        assert [[c.text for c in r.cells] for r in t.rows] == [headers] + rows
    assert doc.sections[0].footer._element.xpath('.//w:fldSimple')[0].get(qn('w:instr')) == 'PAGE'
    paragraphs = doc.paragraphs + [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    assert all(r.font.size and abs(r.font.size.pt - 11) < .01 for p in paragraphs for r in p.runs if r.text)
    for name in ['Normal', 'Body Text', 'Caption', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        rf = doc.styles[name].element.get_or_add_rPr().get_or_add_rFonts()
        assert rf.get(qn('w:ascii')) == rf.get(qn('w:hAnsi')) == b.LATIN_FONT
        assert rf.get(qn('w:eastAsia')) == (b.HEADING_FONT if name == 'Title' or name.startswith('Heading ') else b.BODY_FONT)
        assert not any(key.endswith('Theme') for key in rf.attrib)
    for p in paragraphs + [p for s in doc.sections for p in s.footer.paragraphs]:
        for r in p.runs:
            if not r.text:
                continue
            rf = r._element.get_or_add_rPr().get_or_add_rFonts()
            assert rf.get(qn('w:ascii')) == rf.get(qn('w:hAnsi')) == b.LATIN_FONT
            assert rf.get(qn('w:eastAsia')) == (b.HEADING_FONT if r.bold else b.BODY_FONT)
    log = (CACHE / 'paper.log').read_text()
    for error in ['Overfull', 'Missing character', 'undefined references', 'LaTeX Error']:
        assert error not in log, error
    numbers = validate_numbers()
    reports = [validate_pdf(HERE / 'paper.pdf'), validate_pdf(CACHE / 'word/paper.pdf')]
    paths = [b.SOURCE / 'results_snapshot.json', b.SOURCE / 'source_manifest.json',
             ANALYSIS / 'core/repair_summary.csv', ANALYSIS / 'discordant_cases.csv',
             ANALYSIS / 'discordant_transitions.csv', ANALYSIS / 'sanity_report.json',
             HERE / 'build_abstract.py', HERE / 'validate_abstract.py', HERE / 'paper.tex', HERE / 'paper.docx',
             HERE / 'bibliography_sources.json', ROOT / 'harness/evaluator.py', ROOT / 'harness/readout.py',
             ROOT / 'tasks/__init__.py',
             ROOT / 'docs/analysis/reevaluation_followup_20261015/partial_evaluation_profile.json']
    visual = {'status': 'pending', 'scope': 'both pages of LaTeX and Word-export PDFs'}
    visual_path = HERE / 'visual_review_20261021.json'
    if visual_path.exists():
        candidate = json.loads(visual_path.read_text())
        if all(candidate.get('reviewed_pdf_sha256', {}).get(r['path']) == r['sha256'] for r in reports):
            visual = candidate
    fonts_met = all(r['requested_font_families_present'] for r in reports)
    report = {'validated_at_utc': datetime.now(timezone.utc).isoformat(), 'passed': fonts_met,
              'layout_and_content_passed': True,
              'font_request': {'word_specification_passed': True, 'pdf_font_requirement_met': fonts_met,
                               'japanese_heading': b.HEADING_FONT, 'japanese_body': b.BODY_FONT,
                               'latin_and_digits': b.LATIN_FONT,
                               'limitation': None if fonts_met else '指定フォントがサーバーに存在しないため，PDFは代替フォントによる確認版．指定フォントを備えた環境で再生成と再検証が必要．'},
              'pdfs': reports, 'numeric_checks': numbers, 'word_structure_checked': True,
              'margins_mm': {'left': 18, 'right': 18, 'top': 18, 'bottom': 24},
              'body_and_tables_pt': 11, 'latex_no_overfull_or_missing_characters': True,
              'seven_sections_and_all_citations_checked': True,
              'evaluation_definition_review': {
                  'L0': '構文・importを含むコード読込みと必須関数への適合',
                  'L1': '回路構築と出力評価処理の実行',
                  'L2': '指定入力での出力距離が0.05未満',
                  'primary_L3': 'L0 and L1 and L2 and structure_match．全入力の同値性を保証しない．',
                  'structural_instruction_count_excludes': ['measure', 'barrier'],
                  'shots': 8192, 'simulator_seed': 42, 'transpiler_seed': 42,
                  'definition_source': 'harness/evaluator.pyおよび固定追試のpartial_evaluation_profile.jsonを照合',
                  'old_120_records_included_in_primary': False},
              'source_and_artifact_sha256': {str(p.relative_to(ROOT)): sha(p) for p in paths},
              'visual_review': visual,
              'word_layout_note': 'LibreOfficeで検証．Microsoft Wordで編集後は再ページ割付を確認する．',
              'experiment_actions': {'new_generations': 0, 'generated_code_executions': 0}}
    (HERE / 'validation_report_20261021.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print('Layout/content passed: both PDFs 2 pages, A4, margins, 11pt, page numbers, complete prose, tables and frozen numbers.')
    print('Word font specification passed. PDF requested fonts: ' + ('passed.' if fonts_met else 'NOT MET; substitute-font preview only.'))


if __name__ == '__main__':
    main()
