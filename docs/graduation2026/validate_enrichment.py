"""Validate added catalog, saved examples and prompt publication against evidence."""
from pathlib import Path
from datetime import datetime, timezone
from types import SimpleNamespace
import csv
import hashlib
import json
import re
import sys

from docx import Document
import pymupdf

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
from harness.interventions.prompts import build_initial_prompt, build_intervention_prompt
from harness.refine import diagnose
from tasks import all_tasks


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def normalize(text):
    return re.sub(r'\s', '', text).replace('ﬁ','fi').replace('ﬂ','fl').replace('\u00ad','')


def validate():
    manifest = json.loads((HERE/'enrichment_manifest.json').read_text())
    catalog = list(csv.DictReader((HERE/'intervention_catalog.csv').open()))
    examples = json.loads((HERE/'case_studies.json').read_text())
    prompts = json.loads((HERE/'experiment_prompts.json').read_text())
    doc = Document(HERE/'thesis.docx')
    pdf = pymupdf.open(HERE/'thesis.pdf')
    prose = normalize(''.join('\n'.join(p.get_text().splitlines()[:-1]) for p in pdf))
    word_text = '\n'.join(p.text for p in doc.paragraphs)
    tables = [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
    checks = {
        'source_inputs_unchanged':all(digest(ROOT/p)==h for p,h in manifest['inputs_sha256'].items()),
        'publication_outputs_unchanged':all(digest(HERE/p)==h for p,h in manifest['outputs_sha256'].items()),
        'ten_tables':len(tables)==10,
        'catalog_matches_existing_failure_count_table':True,
        'catalog_matches_requested_rates':[[r[c+'_rate_percent'] for c in ['F','D','B4']] for r in catalog]==[
            ['21.4','26.6','24.7'],['60.3','54.0','61.1'],['36.9','27.7','70.8']],
        'catalog_not_a_selection_rule':normalize('最も高い率の条件を無条件に推奨する規則ではない') in prose,
        'two_saved_examples':len(examples)==2 and {e['cohort'] for e in examples}=={'F_only','D_only'},
        'case_hashes_match_frozen_history':True,
        'case_outcomes_and_excerpts_match_saved_code':True,
        'all_prompt_sources_and_templates_in_word':True,
        'all_prompt_sources_and_templates_in_pdf':True,
        'rendered_repair_templates_match_actual_prompt_builder':True,
        'literal_tasks_match_task_registry':{t['task_id']:t['text'] for t in prompts['tasks']}=={t.id:t.prompt_text for t in all_tasks()},
        'prompt_provenance_limits_explicit':normalize('個別のネットワーク要求全文を直接記録したものではない') in prose,
        'saved_examples_not_causal_claims':normalize('単独の変更や説明指示が成功原因だったとは断定しない') in prose,
        'method_description_replaces_work_log':normalize('本研究の分析には，保存済みの実験結果と必要な部分追試を用いた') in prose
            and all(t not in prose for t in ['今回の執筆','今回の資料作成','今回の原稿修正']),
        'requested_Japanese_revisions':all(t not in prose for t in ['追加成功を得る','分母を伴う分析','F/Dの大小が異なった','未成功を含む消費']),
    }
    for row in catalog:
        old = next(r for r in tables[4][1:] if r[0]==row['label'])
        for c,index in [('F',3),('D',4),('B4',5)]:
            checks['catalog_matches_existing_failure_count_table'] &= int(old[1])==int(row['initial_failures']) and int(old[index])==int(row[c+'_successes']) and f'{100*int(old[index])/int(old[1]):.1f}'==row[c+'_rate_percent']
    history = list(csv.DictReader((ROOT/'docs/analysis/reachable_rescues_20261016/reachable_rounds.csv').open()))
    histories={(r['record_key'],r['round']):r for r in history}
    for example in examples:
        raw={}
        for line in (ROOT/example['source_dir']/'raw.jsonl').open():
            item=json.loads(line)
            if (item['model_spec'],item['task_id'],item['base_seed'])==(example['model_spec'],example['task_id'],example['base_seed']):
                raw[item['condition_id']]={str(a['round']):a['extracted_code'] for a in item['rounds']}
        assert len({r['rounds'][0]['code_sha256'] for r in example['records']})==1
        for record in example['records']:
            for rr in record['rounds']:
                original=histories[(record['record_key'],rr['round'])]
                checks['case_hashes_match_frozen_history'] &= all(original[k]==v for k,v in rr.items())
                checks['case_outcomes_and_excerpts_match_saved_code'] &= hashlib.sha256(raw[record['condition_id']][rr['round']].encode()).hexdigest()==rr['code_sha256']
            checks['case_outcomes_and_excerpts_match_saved_code'] &= record['final_success']==(record['rounds'][-1]['new_L3']=='True')
            if record['excerpt']:
                for which,index in [('before',0),('after',-1)]:
                    code=raw[record['condition_id']][record['rounds'][index]['round']]
                    for line in record['excerpt'][which]:
                        checks['case_outcomes_and_excerpts_match_saved_code'] &= line in code and normalize(line) in prose
    source_texts=[prompts['system'],prompts['preventive_appendix'],*[t['text'].strip() for t in prompts['tasks']],*[t['text'] for t in prompts['repair_templates']]]
    for text in source_texts:
        checks['all_prompt_sources_and_templates_in_word'] &= text in word_text
        checks['all_prompt_sources_and_templates_in_pdf'] &= normalize(text) in prose
    task=SimpleNamespace(prompt_text='〈課題仕様〉')
    ev=SimpleNamespace(L0=False,error_category='〈分類〉',error_message='〈例外〉')
    for item in prompts['repair_templates']:
        actual=build_intervention_prompt(item['condition'],task,'〈前回コード〉',ev)
        actual=actual.replace(diagnose(task,ev),'〈直前の評価から生成した診断〉')
        checks['rendered_repair_templates_match_actual_prompt_builder'] &= actual==item['text']
    checks['P_template_matches_actual_prompt_builder'] = build_initial_prompt('preventive_spec',task)=='〈課題仕様〉\n\n'+prompts['preventive_appendix']
    report={'checked_at_utc':datetime.now(timezone.utc).isoformat(),'passed':all(checks.values()),'checks':checks,
        'pages':len(pdf),'tables':len(doc.tables),'case_groups':len(examples),'task_prompts':len(prompts['tasks']),
        'pdf_sha256':digest(HERE/'thesis.pdf'),'docx_sha256':digest(HERE/'thesis.docx'),
        'new_model_requests':0,'generated_code_executions':0,
        'publication_limits':manifest['publication_scope']}
    (HERE/'enrichment_validation_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':report['passed'],'failed':[k for k,v in checks.items() if not v],'pages':len(pdf),'tables':len(doc.tables)},ensure_ascii=False))
    assert report['passed']


if __name__=='__main__':validate()
