"""Verify expanded thesis prose, frozen arithmetic, references and page previews.

Run after convert_thesis.py and validate_template.py, with the existing document
libraries. This is document verification only; generated programs are never run.
"""
from datetime import datetime, timezone
import csv
import hashlib
import json
from pathlib import Path
import re

from docx import Document
import pymupdf
from PIL import Image, ImageDraw
import build
import template_thesis

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
CACHE=ROOT/'.cache/thesis_expansion_20261028'
CACHE.mkdir(parents=True,exist_ok=True)


def body_stats(path):
    d=Document(path); started=False; ps=[]; chapter=None
    for p in d.paragraphs:
        if p.style.name=='Heading 1':
            chapter=p.text
            if chapter=='はじめに': started=True
            if chapter=='参考文献': break
        if started and not p.style.name.startswith('Heading') and p.text.strip():ps.append(p.text)
    return {'body_paragraphs':len(ps),'body_characters_excluding_tables':sum(len(p) for p in ps),
            'all_paragraph_characters_including_toc':sum(len(p.text) for p in d.paragraphs),'tables':len(d.tables)}


def validate():
    pdf=pymupdf.open(HERE/'thesis.pdf')
    norm=lambda s:re.sub(r'\s','',s).replace('\u00ad','').replace('ﬁ','fi').replace('ﬂ','fl')
    # Remove page-number lines so a paragraph split across pages remains contiguous.
    full=norm(''.join('\n'.join(p.get_text().splitlines()[:-1]) for p in pdf))
    missing=[];expected_count=0
    for _,sections in build.CHAPTERS:
        for title,ps in sections:
            for p in ps:
                expected_count+=1
                if norm(template_thesis.remap(p)) not in full:missing.append({'section':title,'paragraph':p[:90]})
    d=Document(HERE/'thesis.docx'); citations=[]
    for p in d.paragraphs:
        if p.style.name=='Heading 1' and p.text=='参考文献':break
        for n in re.findall(r'\[(\d+)\]',p.text):
            if n not in citations:citations.append(n)
    b4=list(csv.DictReader((build.B4_SOURCE/'summary.csv').open()))
    sums={}
    for grouping in ['model','task','initial_failure']:
        for c,total in [('R',96),('F',183),('D',180),('B4',217)]:
            selected=[r for r in b4 if r['condition']==c and r['generation_cap']=='4' and r['grouping']==grouping]
            assert sum(int(r['successes_from_initial_failure']) for r in selected)==total
            assert sum(int(r['initial_failures']) for r in selected)==545
            sums[f'{c}/{grouping}']=total
    cross={r['success_pattern']:int(r['initial_failure_groups']) for r in csv.DictReader((build.B4_SOURCE/'overlap.csv').open())}
    expected_cross={'none':255,'D_only':21,'F_only':13,'F_D_only':39,'B4_only':67,'B4_D_only':19,'B4_F_only':30,'all_three':101}
    checks={'all_expected_body_paragraphs_in_pdf':not missing,'citation_numbers_in_first_use_order':citations==[str(i) for i in range(1,8)],'B4_three_way_cross_counts':cross==expected_cross and sum(cross.values())==545,'frozen_model_task_failure_totals':True,'fixed_title':norm(build.TITLE) in norm(pdf[0].get_text()),'all_seven_bibliography_entries_in_pdf':all(f'[{i}]' in full for i in range(1,8))}
    before_path=CACHE/'before/thesis.docx'
    before=body_stats(before_path) if before_path.exists() else None
    after=body_stats(HERE/'thesis.docx');after['pdf_pages']=len(pdf)
    if before:
        before['pdf_pages']=len(pymupdf.open(CACHE/'before/thesis.pdf'))
        checks['body_at_least_tripled']=after['body_characters_excluding_tables']>=3*before['body_characters_excluding_tables']
    for i,page in enumerate(pdf,1):
        page.get_pixmap(matrix=pymupdf.Matrix(1.2,1.2)).save(CACHE/f'page_{i:02}.png')
    for start in range(1,len(pdf)+1,6):
        canvas=Image.new('RGB',(1800,1740),'#DCE5EB');draw=ImageDraw.Draw(canvas)
        for k,n in enumerate(range(start,min(start+6,len(pdf)+1))):
            im=Image.open(CACHE/f'page_{n:02}.png').convert('RGB');im.thumbnail((570,808))
            x=15+(k%3)*600;y=30+(k//3)*870
            canvas.paste(im,(x,y));draw.text((x,y-18),f'Page {n}',fill='black')
        canvas.save(CACHE/f'contact_{start:02}.png')
    report={'validated_at_utc':datetime.now(timezone.utc).isoformat(),'passed':all(checks.values()),'checks':checks,'before':before,'after':after,'growth_ratio':after['body_characters_excluding_tables']/before['body_characters_excluding_tables'] if before else None,'expected_body_paragraphs':expected_count,'missing_body_paragraphs':missing,'citation_first_use':citations,'frozen_totals':sums,'pdf_sha256':hashlib.sha256((HERE/'thesis.pdf').read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256((HERE/'thesis.docx').read_bytes()).hexdigest(),'scope':'本文拡充・記述的計算の照合のみ。新規実験・モデル要求・生成コード実行0。'}
    (HERE/'expansion_report_20261028.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False,indent=2));assert report['passed']


if __name__=='__main__':validate()
