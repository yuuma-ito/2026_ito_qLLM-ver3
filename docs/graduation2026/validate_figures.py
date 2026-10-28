"""Verify figures against frozen CSVs and their actual DOCX image parts."""
from pathlib import Path
from datetime import datetime, timezone
import csv
import hashlib
import json
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
import pymupdf
import thesis_figures as figures

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]


def validate():
    m=json.loads((HERE/'figures/figure_manifest.json').read_text())
    doc=Document(HERE/'thesis.docx');pdf=pymupdf.open(HERE/'thesis.pdf')
    norm=lambda s:re.sub(r'\s','',s)
    alltext=norm(''.join(p.get_text() for p in pdf))
    rows=list(csv.DictReader(figures.DATA.open()))
    live=[next(r for r in rows if r['condition']==c and r['grouping']=='overall' and r['generation_cap']=='4') for c in ['R','F','D','B4']]
    count_checks=all(a['condition']==b['condition'] and a['added_successes']==int(b['successes_from_initial_failure']) and a['denominator']==int(b['initial_failures']) and abs(a['rate_percent']-100*int(b['successes_from_initial_failure'])/545)<1e-10 for a,b in zip(m['rates'],live))
    o=next(r for r in csv.DictReader(figures.OVERLAP.open()) if r['scope']=='pooled' and r['grouping']=='overall')
    with zipfile.ZipFile(HERE/'thesis.docx') as z:
        embedded=[hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')]
    texts=[p.text for p in doc.paragraphs];captions=[]
    for p in doc.paragraphs:
        if any('SEQ Figure' in e.text for e in p._p.iter(qn('w:instrText'))):captions.append(p.text)
    checks={
        'frozen_sources_unchanged':all(figures.digest(ROOT/p)==h for p,h in m['sources_sha256'].items()),
        'figure_artifacts_unchanged':all(figures.digest(HERE/'figures'/p)==h for p,h in m['artifacts_sha256'].items()),
        'rates_match_frozen_csv':count_checks,
        'overlap_matches_csv':all(m['fd_overlap'][k]==int(o[k]) for k in ['F_only','both','D_only','neither']),
        'flow_denominators_and_cap':m['flow']=={'initial_groups':1000,'initial_successes':455,'initial_failures':545,'additional_cap':3,'stopping':'L3 success or cap'},
        'all_3_pngs_embedded':len(doc.inline_shapes)==3 and all(m['artifacts_sha256'][s['stem']+'.png'] in embedded for s in m['figures']),
        'captions_in_pdf':len(captions)==3 and all(norm(t) in alltext for t in captions),
        'first_use_definitions_in_summary':all(norm(t) in norm(pdf[1].get_text()) for t in ['モデル・タスク・seedの組合せを初期組と呼び','本研究における最終評価水準（L3','現行評価器の停止・修正規則に従って到達可能な生成履歴']),
        'plain_language_applied':all(t not in ''.join(texts) for t in ['prefix','suffix','追加消費']),
    }
    standalone={}
    for spec in m['figures']:
        f=pymupdf.open(HERE/'figures'/(spec['stem']+'.pdf'))
        standalone[spec['stem']]={'pages':len(f),'fonts_embedded':all(f.extract_font(x[0])[3] for p in f for x in p.get_fonts())}
    checks['standalone_vector_figures']=all(x['pages']==1 and x['fonts_embedded'] for x in standalone.values())
    report={'checked_at_utc':datetime.now(timezone.utc).isoformat(),'passed':all(checks.values()),'checks':checks,'standalone':standalone,'figure_captions':captions,'pages':len(pdf),'pdf_sha256':figures.digest(HERE/'thesis.pdf'),'source_manifest':'figures/figure_manifest.json','new_generations':0,'generated_code_executions':0}
    (HERE/'figure_validation_report_20261028.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':report['passed'],'checks':checks},ensure_ascii=False,indent=2));assert report['passed']


if __name__=='__main__':validate()
