"""Build editable graduation drafts from frozen public aggregates; no experiments.

Dependencies: python-docx, python-pptx, pymupdf (PDF validation only).
Run from repository root: PYTHONPATH=.cache/graduation_tools .venv/bin/python docs/graduation2026/build.py
"""
from pathlib import Path
import csv
import hashlib
import json
from datetime import datetime, timezone
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt as PPt
from pptx.dml.color import RGBColor as PColor
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
SOURCE = ROOT / 'docs/presentation_sources/20261016T_reachable_rescues'
ANALYSIS = ROOT / 'docs/analysis/reachable_rescues_20261016'
SNAP = json.loads((SOURCE/'results_snapshot.json').read_text())
ROWS = list(csv.DictReader((ANALYSIS/'core/repair_summary.csv').open()))
B4_SOURCE = ROOT / 'docs/analysis/independent_4_completed_20261022'
B4_ROWS = list(csv.DictReader((B4_SOURCE/'summary.csv').open()))

def b4_rows(grouping, value='all'):
    return [r for r in B4_ROWS if r['grouping']==grouping and r['group']==value and r['generation_cap']=='4' and r['condition'] in ['B4','R','F','D']]

CONDITIONS = [('R','self_refine'),('F','execution_feedback'),('D','self_debugging')]
MODELS = [('Q4','qwen3.5:4b'),('G4','gemma4:e4b-it-qat'),('Q9','qwen3.5:9b'),('C7','qwen2.5-coder:7b'),('G3','gemma3:4b')]
TITLE = '量子LLM支援プログラミングにおける失敗類型と介入カタログの実証研究'
TITLE_EN = 'An Empirical Study of Failure Types and an Intervention Catalog in Quantum LLM-Assisted Programming'
FONT = 'IPAexMincho'
GOTHIC = 'IPAexGothic'
REFS = [
 'A. Madaan et al.：Self-Refine: Iterative Refinement with Self-Feedback，arXiv:2303.17651v2（2023）．https://arxiv.org/abs/2303.17651v2',
 'X. Chen，M. Lin，N. Schärli，D. Zhou：Teaching Large Language Models to Self-Debug，arXiv:2304.05128v2（2023）．https://arxiv.org/abs/2304.05128v2',
 ('伊藤悠真：現行救済経路の固定発表出典（2026-10-16版）．\n'
  '固定出典：docs/presentation_sources/20261016T_reachable_rescues/\n'
  '参照ファイル：results_snapshot.json，source_manifest.json および対応CSV．\n'
  'B4追加実験（2026-10-22固定）．\n'
  '固定出典：docs/analysis/independent_4_completed_20261022/\n'
  '参照資料：保存設定・manifest・sanity・集計CSV．'),
]

def row(group, value, condition):
    key={'model':'model_spec','task':'task_id','baseline_error':'baseline_error_category'}[group]
    return next(r for r in ROWS if r['scope']=='pooled' and r['grouping']==group and r[key].removeprefix('ollama:')==value and r['condition_id']==condition)

def model_rows(rates=False):
    out=[]
    for short, model in MODELS:
        rs=[row('model',model,c) for _,c in CONDITIONS]
        out.append([short,rs[0]['initial_failures']]+[f"{r['rescued']} ({100*float(r['rescue_rate']):.1f}%)" if rates else r['rescued'] for r in rs])
    return out

ERRORS=[('import失敗','import_error'),('回路構築失敗','build_error'),('構文誤り','syntax'),('出力不一致','wrong_output'),('IF不適合','interface_mismatch'),('量子ビット数','qubit_count_mismatch'),('実行時（未定義変数）','unknown_error')]
def error_rows():
    result=[]
    for label,error in ERRORS:
        rs=[row('baseline_error',error,c) for _,c in CONDITIONS]
        result.append([label,rs[0]['initial_failures']]+[r['rescued'] for r in rs])
    return result

def setfont(run, size=10.5, bold=False, font=FONT):
    run.font.name=font; run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor(0,0,0)
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+key),font)

def field(p, instruction, value='1'):
    r=p.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); r._r.append(begin)
    i=OxmlElement('w:instrText'); i.text=instruction; r._r.append(i)
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); r._r.append(sep)
    p.add_run(value)
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); p.add_run()._r.append(end)

def page_numbers(sec, fmt='decimal', start=1):
    sec.footer.is_linked_to_previous=False
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field(p,'PAGE',str(start))
    el=OxmlElement('w:pgNumType'); el.set(qn('w:fmt'),fmt); el.set(qn('w:start'),str(start)); sec._sectPr.append(el)
    for r in p.runs: setfont(r,10)

def setup(size=10.5, abstract=False):
    d=Document(); s=d.sections[0]; s.page_width=Mm(210); s.page_height=Mm(297)
    s.left_margin=s.right_margin=Mm(18 if abstract else 25)
    s.top_margin=Mm(18 if abstract else 25); s.bottom_margin=Mm(24 if abstract else 25)
    s.header_distance=Mm(10); s.footer_distance=Mm(12)
    for name in ['Normal','Body Text','Caption']:
        st=d.styles[name]; st.font.name=FONT; st.font.size=Pt(size)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
        pf=st.paragraph_format; pf.space_after=Pt(0 if abstract else 6); pf.line_spacing=Pt(12.7 if abstract else 17)
        pf.widow_control=True
    for name,pt in [('Title',18),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        st=d.styles[name]; st.font.name=GOTHIC; st.font.size=Pt(pt); st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),GOTHIC)
        st.paragraph_format.space_before=Pt(8); st.paragraph_format.space_after=Pt(5)
    d.core_properties.author='伊藤 悠真'; d.core_properties.title=TITLE; d.core_properties.subject='卒業研究発表会用草稿・2026-10-16固定集計'
    d.core_properties.comments='検証済み保存結果に基づく草稿．提出年月・所属英訳・教員確認・類似レポートは要確認．'
    return d

def para(d,text,size=None,bold=False,center=False,indent=True):
    p=d.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0 if center or not indent else size or 10.5)
    if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if size is not None: setfont(p.add_run(text),size,bold)
    else: p.add_run(text)
    return p

def heading(d,text,level=2,small=False):
    p=d.add_paragraph(text,style='Heading '+str(level))
    if small:
        p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(3)
        for r in p.runs:setfont(r,10.5,True,GOTHIC)
    return p

def table(d,caption,headers,rows,size=10.5,widths=None):
    p=para(d,caption,size,center=True,indent=False);p.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    cols=d.sections[-1]._sectPr.find(qn('w:cols'))
    narrow=cols is not None and cols.get(qn('w:num'))=='2'
    width=81 if narrow else 160
    ratios=([w/sum(widths) for w in widths] if widths else ([.32,.17,.17,.17,.17] if narrow and len(headers)==5 else [1/len(headers)]*len(headers)))
    for col,ratio in zip(t.columns,ratios):col.width=Mm(width*ratio)
    tw=t._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(width/25.4*1440)))
    for c,text in zip(t.rows[0].cells,headers): c.text=str(text)
    for rowdata in rows:
        for c,text in zip(t.add_row().cells,rowdata):c.text=str(text)
    for i,r in enumerate(t.rows):
        trPr=r._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        if i==0:trPr.append(OxmlElement('w:tblHeader'))
        for c,ratio in zip(r.cells,ratios):
            c.width=Mm(width*ratio)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=Pt(size+2)
                for run in p.runs:setfont(run,size,i==0)
            tcPr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            if i in [0,len(t.rows)-1]:
                for edge in (['top','bottom'] if i==0 else ['bottom']):
                    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4');borders.append(e)
            tcPr.append(borders)
    end=d.add_paragraph();end.paragraph_format.space_after=Pt(0);end.paragraph_format.line_spacing=Pt(2)
    setfont(end.add_run(''),2)
    return t

JA_SUMMARY='''LLMによる量子回路コード生成では，APIの誤用や回路構築の失敗に加え，実行できても期待した出力を返さない失敗が生じる．本研究では，Qiskitを用いた量子回路コード生成において，初期失敗に対する修正介入と独立再生成を比較し，追加成功数，成功対象および計算コストを分析した．自己見直しR，実行結果提示F，説明・原因分析を促す自己デバッグDと，独立再生成B4を，同じ初期コードを起点として比較した．生成前注意事項Pは独立生成として扱った．全1,000初期組の最終L3成功率はB45.5%，P47.3%，R55.1%，F63.8%，D63.5%，B4 67.2%であった．
対象は5モデル，10タスク，seed 0〜19の1,000初期組，5,000条件記録である．検証済み保存prefixと必要な部分追試suffixを組み合わせた到達可能な経路を採用した．成功は指定入力の出力一致と所定の構造要件を含むL3で判定した．初期失敗545組の救済はR96組（17.6%），F183組（33.6%），D180組（33.0%）であった．FとDの共通救済は140組，Fのみ43組，Dのみ40組であり，救済数が近くても対象は異なった．Fのみ・Dのみの救済計83組について，両条件166履歴・398修正遷移を静的に確認した．
独立再生成B4では同じ初期コードを保持し，初期失敗に対して最大3回，前回コードや診断を与えず再生成した．B4は217/545組（39.8%）で成功し，最終成功は672/1,000組（67.2%）であった．追加1,342要求，995,815 tokens，追加tokens／追加成功は4,589.0であった．最大生成回数が同じでも入力情報・seed規則・実行時期が異なり，因果効果は断定しない．
未救済例への消費を含む追加tokens／救済はR18,852.6，F9,333.8，D10,976.6であった．3回目の修正はFで28組，Dで33組を追加救済した．これらは保存記録と部分追試を組み合わせた計算値であり，全件新実験の実測コストではない．結果は今回のモデル・タスクと指定入力での評価に限られ，介入の一般的優劣，因果効果，統計的有意差を主張しない．今後は介入選択・停止方策を事前固定し，別seed・タスクで検証する必要がある．'''
EN_SUMMARY='''Large language models can produce Qiskit code that fails during import or circuit construction, or produces an unexpected output despite executing successfully. This study examines how often repair interventions rescue initial failures and whether different interventions rescue the same cases. We compare self-review (R), execution feedback (F), and self-debugging with explanation and cause analysis (D), starting from shared initial code. Preventive instructions (P) are treated as independent generation.
The analysis covers 1,000 initial groups and 5,000 condition records across five models, ten tasks, and seeds 0–19. The primary paths combine verified saved prefixes with necessary suffixes from partial reruns. Success requires the expected output for the specified input and predefined structural constraints. Of 545 initial failures, R rescued 96 (17.6%), F rescued 183 (33.6%), and D rescued 180 (33.0%). F and D shared 140 rescued cases; 43 were rescued only by F and 40 only by D. All 166 condition histories and 398 repair transitions for these 83 discordant cases were inspected statically.
An additional independent-regeneration control (B4) retained the shared initial code and generated up to three new candidates from the task specification without previous code or diagnostic feedback. B4 succeeded on 217 of 545 initial failures (39.8%), yielding 672 successes out of 1,000 groups (67.2%). It used 1,342 additional requests and 995,815 tokens, or 4,589.0 tokens per added success. Equal maximum generation counts do not equalize input information, seed rules, or execution periods.
Additional input and output tokens per rescue, including consumption on unresolved cases, were 18,852.6 for R, 9,333.8 for F, and 10,976.6 for D. The third repair added 28 rescues for F and 33 for D. These costs describe reconstructed paths, rather than a fully regenerated experiment. The findings are limited to the selected models, tasks, inputs, and evaluation policy. They do not establish general superiority, causal effects, or statistical significance. Future work should predefine intervention selection and stopping policies and evaluate them on separate seeds and tasks.'''

# Each chapter has substantive prose and tables; no invented experiments or results.
CHAPTERS=[
('はじめに',[
('研究背景',[
'量子回路をプログラムとして構築する際には，量子状態や測定分布に関する理解と，ソフトウェアAPIの正しい使用の両方が必要になる．LLMの生成コードを利用する場合も，構文上の正しさだけで有用性を判断できない．importが成功しても回路構築に失敗する場合があり，実行が完了しても期待した出力と一致しない場合がある．',
'生成を一度で終えず，前回コードを見直させたり，実行結果を提示したりする方法が考えられる．自己フィードバックによる反復改善のSelf-Refine [1]と，コードの説明や実行結果を用いたSelf-Debugging [2]は，こうした修正の背景となる．ただし，一般的なコード生成での知見が，量子回路コードの各種失敗に同じ形で当てはまるとは限らない．',
'実行結果には，存在しないAPIの使用を知らせる診断と，期待出力との不一致を知らせる診断が含まれる．両者では修正に必要な情報が異なる可能性がある．全体の最終成功率だけを比較すると，初期成功をそのまま保持した分と，修正によって救済した分が混在する．そこで本研究では，初期失敗を分母に固定して介入を比較する．']),
('研究目的と研究課題',[
'本研究の目的は，Qiskit量子回路コードの初期失敗に対する修正介入と独立再生成の有効性を比較し，失敗類型やモデルに応じた対処方法を選択するための根拠を整理することである．第一の問いは，同じ初期失敗に対して修正介入R/F/Dと独立再生成B4がそれぞれどれだけ追加成功を得るかである．第二の問いは，救済数が近い介入が同じ失敗を救済しているかである．第三の問いは，未救済例も含む追加消費と修正回数が救済にどう対応するかである．',
'本研究の貢献は，共有初期コードを用いた修正介入と独立再生成の追加成功率の比較，初期失敗類型・モデル・タスク別の分母を伴う分析，FとDの片側救済全83組の履歴整理，および追加修正コストと3回目の寄与の記述である．介入の最適な選択方策を既に実証したとは主張しない．']),
('論文の構成と対象範囲',[
'第2章で関連研究と本研究の位置づけを述べ，第3章で介入と指標を定義する．第4章では評価・記録の実装を説明し，第5章で設定と出典を整理する．第6章で結果，第7章で考察と限界，第8章で結論と今後の課題を述べる．本研究の分析には，保存済みの実験結果と必要な部分追試を用いた．B/R/F/D/Pは2026-10-16版，B4は2026-10-22固定の追加実験結果 [3] に基づく．'])]),
('関連研究と本研究の位置づけ',[
('自己フィードバックによる修正',[
'Self-Refine [1] は，同じLLMによる生成，フィードバック，修正の反復を扱う．本研究のRは，前回コードと課題仕様を与えて自己点検を求める条件であり，外部評価の診断内容を与えない．自己点検だけで救済できる範囲を確認する比較条件として用いる．',
'原論文と本実装ではプロンプト構成，対象課題，出力制約，モデルが異なる．したがって本研究のRの結果を，原論文の手法全体の性能と同一視しない．先行研究は設計の背景であり，完全再現実験として扱わない．']),
('実行結果と自己デバッグ',[
'Self-Debugging [2] は，生成コードの説明や実行結果の検討を用いた自己デバッグを扱う．本研究ではFを実行ベースの診断を用いる条件，Dを同じ方式の診断に加えて説明・原因分析・修正の順序を促す条件として実装した．Dの追加指示と出力長の違いを含めて観測する．',
'Dは説明と原因分析を修正版コード先頭のコメントに残す実装である．モデルが内部でどのように推論したかは測定していない．Dで救済されたという事実だけから，説明が正しかった，あるいは説明そのものが成功原因だったとは判断できない．']),
('量子回路コードで確認する対象',[
'本研究は，回路を返す関数の存在，回路の構築と評価処理，指定入力の期待状態または分布，所定の構造要件を区別する．量子回路では，実行可能性と期待出力への適合は別の性質である．測定分布が一致しても，全入力で同じアルゴリズムを実装している保証にはならない．',
'本研究は，選定した量子回路課題における介入の救済対象と追加計算コストを記述する実証研究である．一般の量子プログラム全体への網羅的評価や，形式的な同値性証明は対象に含めない．その範囲を明確にしたうえで，成功例と失敗例を同じ分母で扱う．'])]),
('介入条件と評価指標',[
('6条件と共有初期コード',[
'表3.1に6条件の方法，最大生成回数と役割を示す．モデル・タスク・seedの組を初期組と呼ぶ．Bは初期生成のみの条件である．R/F/DはBと同じ初期コードと判定を共有し，初期失敗に対して最大3回の修正を行う．成功した時点で停止し，以降の修正は行わない．R/F/D/B4の最大4回は，初回1回と最大3回の追加生成を意味する．B4もBの初期コードを共有する一方，PはBとは異なる初期コードを独立に生成する．',
'Rは自己見直し，Fは実行結果提示，Dは説明・原因分析を促す自己デバッグである．Pは初回プロンプトに注意事項を追加して独立に生成する条件であり，同じコードを修正する条件ではない．Pの対応比較は，Bの初期失敗に対応して成功した場合と，Bの初期成功に対応して失敗した場合を分ける．']),
('独立再生成対照B4',[
'B4は初期1回を含む最大4回生成である．Bと同じ初期コード・L3判定を用い，初期成功455組は保持する．初期失敗545組のみ，前回コードと診断を与えず課題仕様から最大3回独立再生成し，L3成功時に停止する．初期失敗からの成功は修正介入による救済と区別する．',
'5モデル・10タスク・base seed 0〜19は従来と対応するが，追加生成はindependent_4_v1の専用seed規則を用いる．temperature 0.7，max_tokens 1,024，timeout 600秒，並列1，reasoning_effort noneで実行した．上限が同じでも入力情報・実要求数・tokens・実行時期・診断文・出典schemaは等しくない．']),
('L3と出力一致',[
'L0は生成コードの読込み・import等と必須インターフェースの要件，L1は回路構築と評価処理の実行，L2は指定入力に対する期待出力との一致を表す．L3はL0，L1，L2およびstructure_matchがすべて真であることと定義する．構造要件には量子ビット数と測定・barrierを除く回路命令数の範囲を用いる．',
'測定分布の距離は全変動距離 TVD(p,q) = (1/2)Σx |p(x)−q(x)| とし，8,192 shotsで得た分布と期待分布を比較する．状態ベクトルは大域位相を整合したユークリッド距離を用いる．各タスクで距離0.05未満を出力一致とする．測定・変換seedは42に固定したが，有限shotsと指定入力に限られる．',
'本稿で用いるevaluation_policy_v4は，測定先から量子ビット順への対応を反映する．タスクに構造制約が明示されていない項目を，一律の制約として追加しない．L3は所定の検査を通過したことを意味し，全入力でのアルゴリズム同値性やプロンプトの全自然言語条件の充足を保証しない．']),
('救済・保持・独立生成の比較',[
'R/F/Dによる修正で，初期L3未達から最終L3成功に変化した初期組を救済と呼ぶ．救済率 = 救済組数／初期L3未達組数 とする．初期失敗類型は介入前のBの分類に固定する．修正後の分類で分母を選別すると対象が変わるため，その方法は採用しない．',
'R/F/Dの最終成功には初期成功の無修正保持が含まれる．成功時停止により退行0となる設計であり，退行0をモデルが常に正しさを維持できる証拠と解釈しない．Pでは独立生成による改善と悪化の両方を報告する．'])]),
('実装と記録の検証',[
('ハーネスの構成',[
'実装はPythonを中心とし，tasks/が課題仕様・期待値・構造条件を，harness/interventions/prompts.pyが条件別の修正指示を保持する．harness/llm_clients.pyがモデル要求を扱い，harness/experiment.py等が初期共有と反復を管理する．harness/evaluator.pyとharness/readout.pyが回路と出力を検証する．これらの責務を分け，生成の失敗と評価の不具合を区別して追跡する．',
'評価対象コードはbuild_circuit()を定義し，QuantumCircuitを返すインターフェースを前提とする．測定分布と状態ベクトルのどちらを検査するかはタスク仕様で定める．実装環境の依存範囲はrequirements.txtに保持し，過去の実験の実行環境を現在の環境から推定して埋めない．保存設定・ハーネス版・モデルdigestは元実験manifestを参照する．']),
('失敗分類とモデル接続障害',[
'構文，import，必須インターフェース，回路構築，出力不一致，量子ビット数などの分類を記録する．コード側のimport_errorやbuild_errorは介入の対象となる実験結果である．一方，モデルAPIの接続エラーやタイムアウトは生成されたコードの品質と分けて検証する．',
'保存時にunknown_errorと記録された初期失敗1組について，事後確認により未定義変数simによるNameErrorと判明したため，本稿では「実行時エラー（未定義変数）」と表記する．保存データの分類キーは変更していない．API接続障害には含めない．固定出典の検証状態を参照し，未解決の実行・評価問題を確定結果のように扱わない．']),
('到達可能な履歴の再構成',[
'評価方針の修正は，停止する時点や次に与える診断文を変えることがある．保存コードを新しい評価で再採点するだけでは，新診断を受けて生成されるはずの後続コードを再現できない．そこで，検証済みの生成履歴の前半部分と，必要な部分追試の後半履歴（suffix）を接続した現在の到達経路を主結果とする．',
'部分追試は293ケースで，部分追試で生成した後半履歴は647ラウンドである．分岐後に到達できない旧保存suffixを現在の経路に流用しない．初期生成から全件を新評価で生成し直した実験とは区別する．現行の到達roundは9,249であり，5,000条件記録とは集計単位が異なる．']),
('集計の整合性と静的履歴分析',[
'固定集計では全5,000条件の対応，欠落・重複・予定外記録，共有初期の一致，成功時停止，層別分母の合計，出典と到達コードのhashを検査した．固定sanity reportは成功しており，欠落・重複・予定外記録・共有初期不一致は0である．本稿の表は公開CSVから取得し，旧値を転記して混在させない．',
'片側救済83組については両条件166履歴，398修正遷移をコード文字列として静的確認した．import/API，回路構築・インターフェース，量子ゲート等のテキスト，測定・出力取得を複数該当可で注記した．生成コードを実行しない分析であり，変更の意味や因果的寄与は別の検証を要する．'])]),
('実験設定と出典',[
('対象モデルと課題',[
'対象はQ4，G4，Q9，C7，G3の5モデルである．正式なモデルタグは表5.1に示す．タグ中の数値を厳密なパラメータ数とみなさず，モデル識別子として扱う．各モデル200初期組であり，10タスクとseed 0〜19を含む．',
'タスクはBell，GHZ，Deutsch–Jozsa constant/balanced，Bernstein–Vazirani，Grover，QFT，IQFT，QPE，Ansatzの10種類である．各タスクは指定入力・期待値を持つ．タスク名が同じでも任意入力を網羅する評価ではない．ハーネスの難度分類を利用するが，普遍的な難しさを表す尺度として解釈しない．']),
('生成条件と元実験の区別',[
'temperatureは0.7，1要求の生成上限は1,024 tokens，最大修正回数は3回である．保存設定ではreasoning_effortはnone，timeout_secは600，parallel_requestsは1である．実際に生成された長さは可変であり，上限の同一性は消費量の同一性を意味しない．',
'元実験は，3モデルseed 0〜9の1,500条件記録，5モデルseed 10〜19の2,500条件記録，残り2モデルseed 0〜9の1,000条件記録に分かれる．schemaや診断文処理の差を保持したうえで，初期組と条件を対応づける．元実験IDと保存設定の参照先は付録Aに示す．',
'旧120件は2モデル，seed 0〜2，2条件，L2終了の既報であり，現行のL3による5,000条件記録へ合算しない．旧評価，保存コードの再採点，部分追試を三つの独立した主結果として並べない．R/F/D/PおよびBの主集計には2026年10月16日版の固定スナップショットを用い，B4については2026年10月22日版の追加実験結果を用いる [3]．']),
('追加コストと修正回数の集計方針',[
'追加修正コストは共有初期生成を除き，失敗した修正への消費を含めた入力＋出力tokensの合計である．tokens／救済は総追加tokensを救済組数で割る．要求時間は保存されたelapsed_secの合計であり，評価時間や全工程の所要時間ではない．',
'保存履歴の前半の記録値と，部分追試による後半の実測値を分けて保持する．合計は再構成した経路に対する計算値である．金銭単価を測定していないため，金銭費用の削減とは呼ばない．修正上限別の集計も同じ最大3修正の観測履歴を区切ったものであり，上限ごとに独立実験した結果ではない．'])]),
('結果',[
('全体の救済率',[
'1,000初期組中，初期成功455組，初期失敗545組であった．表6.1に条件別の救済を示す．Rは96組（17.6%），Fは183組（33.6%），Dは180組（33.0%）を救済した．FとDは本対象の初期失敗のおよそ3分の1を救済し，Rを上回った．全体率は545組を分母とする合算値であり，モデル率の単純平均ではない．',
'R/F/Dの最終成功は551，638，635組で，初期成功455組の無修正保持を含む．FとDの最終成功差は3組だが，これは同じ対象を救済したことを意味しない．集計単位と分母をそろえ，件数と率を併記して解釈する．']),
('モデル別・初期失敗類型別の結果',[
'表6.2のモデル別集計では，G4でF48/62組，D38/62組に対し，Q4ではF78/162組，D90/162組と大小が入れ替わる．C7は両条件23/65組，G3は両条件1/190組であった．モデルごとの初期失敗数は異なるため，救済数だけを比較しない．G3の190失敗組も全体分母に含める．',
'表6.3ではimport失敗に対してF58/271組（21.4%），D72/271組（26.6%），回路構築失敗に対してF76/126組（60.3%），D68/126組（54.0%）であった．出力不一致はF24/65組，D18/65組である．初期分類とモデル・タスクの関連があるため，この表だけから最適な介入選択を決めない．少数分母の率は参考値として扱う．']),
('タスク・難度別の結果',[
'表6.4に全10タスクの分母と件数を示す．BellはF10/34組・D12/34組，GHZは12/35組・13/35組，QFTは34/67組・34/67組，QPEは13/94組・11/94組であった．同じ救済数でも分母が異なる場合がある．また同じタスク内にモデルとseedの違いが含まれる．',
'ハーネス難度別では易がF22/69組・D25/69組，中が79/225組・79/225組，難が82/251組・76/251組であった．難度間の差にはタスク構成の違いが含まれる．モデル×タスクや元実験別・seedブロック別の分母は固定CSVに保持しており，集約値だけを根拠に一般化しない．']),
('FとDの救済対象の重なり',[
'表6.5のとおり，FとDの共通救済は140組，Fのみ43組，Dのみ40組，双方未救済322組であった．片側だけで救済された組は合計83組である．少なくとも一方が救済した223/545組（40.9%）は，両条件の結果を見て作った集合である．実行前の選択方策の成功率として扱わない．',
'全83組の両履歴を確認したところ，初期import失敗はFのみ8組・Dのみ22組，回路構築失敗は19組・11組であった．救済側の履歴でimport/API関係の変更がFのみ28/83組・Dのみ34/83組，回路構築・インターフェース関係が24/83組・32/83組，量子ゲート等のテキスト変更が38/83組・32/83組，測定・出力関係が11/83組・16/83組に認められた．重複を許す全83組に対する注記である．',
'静的変更は失敗側にも現れ得る．同じ遷移で複数箇所が変わるため，特定の変更を成功原因と断定しない．全例を除外なく記録し，該当なしや最後まで失敗した履歴もCSVに含めた．']),
('追加消費と3回目の修正',[
'表6.6の追加tokens／救済はR18,852.6，F9,333.8，D10,976.6であった．総追加tokensはR1,809,845，F1,708,086，D1,975,782である．そのうち最終的な未救済例への消費はR1,618,684，F1,374,240，D1,605,388 tokensを占める．成功例だけに消費が発生するわけではない．',
'round別救済はRが57，20，19組，Fが112，43，28組，Dが108，39，33組であった．表6.7に3回目の分母と消費を示す．Fは2回後未救済390組から28組（7.2%），Dは398組から33組（8.3%）を追加救済した．2回で区切ると，3回までに救済できたうちF15.3%，D18.3%を失う．',
'3回目の消費はF491,676，D588,970 tokensであり，最後まで失敗した例への消費も含む．救済を失う割合と追加消費を併記することで，対象範囲で停止方策を検討する材料になる．一般的に2回か3回のどちらが最適かは，この結果だけでは決められない．']),
('Pの独立生成としての比較',[
'PはBの初期失敗に対応する成功124/545組（22.8%）と，初期成功に対応する失敗106/455組（23.3%）を生じ，最終成功は473/1,000組であった．Bとの差18組には改善と悪化の双方が含まれる．R/F/Dと異なり，同じ初期コードの修正効果を表す指標ではない．',
'Pの1,000独立生成の記録消費は764,083 tokensである．これは修正型の共有初期を除く追加消費と集計範囲が異なる．Pの結果を修正型の救済率に単純に並べて，費用効率の優劣を結論づけない．'])]),
('考察と限界',[
('介入で何が変わったか',[
'本対象では，実行結果を使うFとDが，外部診断を与えないRより多くの初期失敗を救済した．これは，課題仕様だけでなく実行に基づく診断を使う条件が有用だったという観測である．ただしFとDにはプロンプト構成や出力消費の違いもあり，情報の各要素の寄与を分離した実験ではない．',
'FとDの全体差3組に対し片側救済は83組であった．全体率の近さは，救済する対象の同一性を意味しない．G4とQ4で大小が入れ替わり，import失敗と回路構築失敗でも傾向が異なる．介入を選ぶには，全体率に加え，対象の失敗構成と適用範囲を確認する必要がある．']),
('介入カタログへ向けた解釈',[
'初期失敗類型別の結果は，介入カタログの候補を作る根拠になる．例えばimport失敗と回路構築失敗ではF/Dの観測された救済率が異なる．しかし，同じ類型でもモデルやタスクの構成が異なり，類型だけを入力にした選択規則の有効性は未検証である．',
'F/Dの和集合223組は事後的に得られる上限的な参照集合である．実行時に正解の介入を事前に選べるわけではない．両方を実行する方策を検討する場合も，追加消費や停止条件を含めた独立の評価が必要になる．']),
('内的妥当性と再構成経路',[
'共有初期コードによって初期生成の違いを抑えている一方，同一タスク・seedの依存が残る．5,000条件記録を独立な5,000標本として扱わない．元実験にはschemaや診断文処理の差もあるため，差を消して単一の同条件実験だったと表現しない．',
'現行経路は検証済み保存prefixと必要な新suffixの組合せである．停止やフィードバックの影響を受ける部分は追試しているが，全件の新生成とは異なる．古い要求ログが完全には残らない出典もあり，負荷・待機・再送の影響を完全には分離できない．コストはこの再構成経路の記録に限った計算値である．']),
('評価範囲・一般化・静的分析の限界',[
'L3は指定入力の出力と所定の構造要件に限定される．全入力の正しさ，すべての自然言語指示，等価なアルゴリズム実装を保証するものではない．有限shots・指定入力による評価範囲の制約，生成長上限による出力制約，モデルdigest・量子化・サーバ負荷の影響も考慮する必要がある．',
'片側救済83組を全件確認しても，静的差分だけで変更の意味的寄与は特定できない．ゲート名の増減は量子論理の訂正の証拠にはならず，複数同時変更がある．さらに，今回の5モデルと10タスクから，モデルの一般的優劣や介入の普遍的優位を断定しない．本研究では因果効果や統計的有意差を主張しない．'])]),
('おわりに',[
('研究成果',[
'本研究では，LLMが生成したQiskit量子回路コードの初期失敗を，修正介入がどれだけ，どの対象で救済するかを分析した．1,000初期組のうち失敗545組に対し，R96組，F183組，D180組を救済した．今回の対象ではFとDがRより多くを救済した．',
'FとDの救済数は近かったが，共通140組，Fのみ43組，Dのみ40組であり，片側救済83組を認めた．全83組の166履歴と398遷移を静的に整理した．追加消費は未救済例も含め，3回目の追加救済と停止回数を減らした際の損失を明示した．']),
('今後の課題',[
'第一に，介入の選び方と停止条件を事前固定し，別seed・タスクで検証する．その際，事後的に最も良い結果を選ぶ評価と区別し，追加要求数と消費を含めて報告する必要がある．第二に，静的変更の意味を人手で確認し，複数同時変更の寄与を分ける検証方法を設計する．',
'第三に，測定乱数，生成長の上限，モデルdigestや実行負荷への感度を調べる．L3を拡張する場合は，過去の成功時停止とフィードバックへの影響を監査する．同じ保存設定と出典を保持し，必要な追試を明示的に計画する．これらは今後の検証課題である．']),
('提出に向けた確認',[
'本稿は検証済みの固定集計に基づく卒業論文草稿である．提出年月，英語の著者表記・所属，指導教員による内容・参考文献書式の確認を残す．要旨の審査用と最終版の提出時には，学部の指示に従って剽窃チェックと類似レポートの作成を行う．本稿の作成をもって剽窃チェック完了とはしない．'])]),
]


# B4 is a separate completed experiment, sharing the initial cohort only.
CHAPTERS[5][1].append(('独立再生成B4との比較',[
'B4実験 independent_4_full_20261020T185811381595Z は2026年10月21日に全1,000組を完了した．欠落・重複・予定外記録はいずれも0で，sanity検証は成功した．追加roundでは未定義変数による実行時エラーが2件あり，いずれも生成コード内のNameErrorであり，API接続障害ではない．',
'B4は初期失敗545組中217組（39.8%）で成功し，最終成功は672/1,000組（67.2%）であった．表6.8，表6.9，表6.10に全体・モデル・初期失敗類型別の比較を示す．R96組（17.6%），F183組（33.6%），D180組（33.0%）に対し，B4の成功数はFより34組，Dより37組多かった．この観測により，修正条件と単回生成だけの比較から介入自体の効果を論じることはできない．',
'B4とFの共通成功は131組，B4のみ86組，Fのみ52組，双方未成功276組であった．B4とDでは共通120組，B4のみ97組，Dのみ60組，双方未成功268組であった．独立再生成が多く成功しても，修正介入だけが成功する対象も残る．事後的な和集合を事前の選択成功率とは扱わない．',
'B4の追加要求は1,342回，追加tokensは995,815，追加tokens／追加成功は4,589.0であった．共有初期は一度だけ数え，追加消費から除外し，最後まで未成功の例への消費を含めた．R/F/Dは過去の保存prefixと必要なsuffixの再構成値であり，B4との要求時間・tokens差を同時期の同一入力条件の実測費用差とは解釈しない．']))
CHAPTERS[6][1].insert(0,('再生成対照を含む解釈',[
'本対象ではB4の初期失敗からの成功率39.8%がF33.6%，D33.0%を上回った．したがって，実行結果を与える修正が独立再生成より常に有用であるとは結論づけられない．一方でFのみ52組，Dのみ60組があり，再生成だけで全介入の成功を包含しない．介入カタログには独立再生成，対象分母，入力情報，追加消費も記載する．',
'B4は初回込み最大4回という生成機会の上限を合わせた比較である．成功時停止による実要求数，入力tokens，専用seed規則，実行時期，保存・追試経路の違いが残る．診断情報の因果効果や統計的有意差をこの集計から認定しない．次は同時期・同一評価設定で方策を事前固定し，別seed・タスクで比較する必要がある．']))
CHAPTERS[7][1][0][1].append('独立再生成B4は同じ初期失敗545組中217組（39.8%）で成功し，最終成功は672組となった．F/Dのみが成功する組も残り，独立再生成を含む選択・停止方策の検証を今後の課題とする．')


# Expand thesis prose without changing the frozen experimental results.
from thesis_expansion import apply as expand_thesis
expand_thesis(CHAPTERS, REFS)
from thesis_enrichment import ensure as ensure_enrichment, apply as enrich_thesis
ensure_enrichment()
enrich_thesis(CHAPTERS)

# Use the same plain-language vocabulary in the abstract and generated body.
from thesis_language import polish, DEFINITIONS
JA_SUMMARY = polish(JA_SUMMARY)
JA_SUMMARY = JA_SUMMARY.replace('対象は5モデル，10タスク，seed 0〜19の1,000初期組，5,000条件記録である．', '対象は5モデル，10タスク，seed 0〜19である．モデル・タスク・seedの組合せを初期組と呼び，B/R/F/D/Pの1,000初期組・5,000条件記録を対象とし，B4の追加実験を別に集計する．')
JA_SUMMARY = JA_SUMMARY.replace('成功は指定入力の出力一致と所定の構造要件を含むL3で判定した．', '本研究における最終評価水準（L3）は，指定入力の出力一致と所定の構造要件を含む成功基準である．')
JA_SUMMARY = JA_SUMMARY.replace('組み合わせた到達履歴を採用した．', '組み合わせ，現行評価器の停止・修正規則に従って到達可能な生成履歴（以下，到達履歴）を採用した．')
CHAPTERS[0][1][0][1].append(DEFINITIONS)
for _chapter_index, (_chapter_title, _sections) in enumerate(CHAPTERS):
    CHAPTERS[_chapter_index] = (polish(_chapter_title), [(polish(t), [polish(p) for p in ps]) for t,ps in _sections])
CHAPTERS[5][1][0][1].insert(2, '図3は同じ初期失敗545組を分母として追加成功率を比較したものである．R/F/Dは修正による救済，B4は独立再生成による追加成功として区別する．')
CHAPTERS[5][1][4][1].insert(1, '図4はこの対応を件数に比例する帯で示したものである．共通救済と片側のみの救済を，双方未救済の集合と合わせて比較できる．')



def abstract():
    # The current university abstract has one canonical prose/table source.
    import runpy
    import shutil
    runpy.run_path(str(ROOT/'docs/midterm2026/build_abstract.py'), run_name='__main__')
    shutil.copy2(ROOT/'docs/midterm2026/paper.docx', HERE/'abstract.docx')


def thesis_tables(d, chapter):
    if chapter == 3:
        from six_conditions import DEFINITIONS
        table(d,'表3.1　6条件の方法・最大生成回数・役割',
              ['条件','方法','最大生成回数','役割'],DEFINITIONS,widths=[16,48,30,66])
    elif chapter == 5:
        table(d,'表5.1　モデルの略号と識別子',['略号','モデルタグ','初期組数'],[[s,m,200] for s,m in MODELS],widths=[18,112,30])
    elif chapter == 6:
        overall={r['condition']:r for r in b4_rows('overall')}
        table(d,'表6.1　初期失敗545組からの追加成功',['条件','追加成功','率','最終成功／1,000'],[
            [c,overall[c]['successes_from_initial_failure'],f"{100*float(overall[c]['success_from_initial_failure_rate']):.1f}%",overall[c]['final_L3_successes']]
            for c in ['R','F','D','B4']],widths=[18,36,36,70])
        for number,grouping,labels,title,widths in [
            (2,'model',[(s,'ollama:'+m) for s,m in MODELS],'モデル別',[35,25,25,25,25,25]),
            (3,'initial_failure',THESIS_ERRORS,'初期失敗類型別',[56,18,21.5,21.5,21.5,21.5]),
            (4,'task',THESIS_TASKS,'タスク別',[48,24,22,22,22,22])]:
            data=[]
            for label,value in labels:
                rs={r['condition']:r for r in b4_rows(grouping,value)}
                data.append([label,rs['B4']['initial_failures']]+[rs[c]['successes_from_initial_failure'] for c in ['R','F','D','B4']])
            table(d,f'表6.{number}　{title}の追加成功数',['対象','分母','R','F','D','B4'],data,widths=widths)
        table(d,'表6.5　F/Dの救済対象（初期失敗545組）',['共通','Fのみ','Dのみ','双方未救済'],[[140,43,40,322]])
        table(d,'表6.6　共有初期を除く追加消費（未成功への消費を含む）',['条件','追加要求','追加tokens','tokens／追加成功'],[
            [c,overall[c]['additional_requests'],f"{int(overall[c]['additional_tokens']):,}",f"{float(overall[c]['tokens_per_added_success']):,.1f}"] for c in ['R','F','D','B4']],widths=[18,30,54,58])
        third={r['condition_id']:r for r in SNAP['third_round']}
        table(d,'表6.7　3回目の追加救済と消費',['条件','追加／未救済','追加tokens','要求秒','2回で失う割合'],[
            [s,f"{third[c]['additional_rescues']}/{third[c]['at_risk_after_2']}",f"{third[c]['third_round_tokens']:,}",f"{third[c]['third_round_seconds']:,.1f}",f"{third[c]['loss_if_stop_at_2']*100:.1f}%"] for s,c in CONDITIONS],widths=[16,34,34,32,44])
    elif chapter == 7:
        from thesis_enrichment import catalog_rows
        table(d,'表7.1　介入カタログの抜粋（主要3類型の観測率）',
              ['初期失敗類型','初期失敗数','F救済率','D救済率','B4追加成功率'],
              catalog_rows(),widths=[40,28,28,28,36])


THESIS_ERRORS=[('import失敗','import_error'),('回路構築失敗','build_error'),('構文誤り','syntax'),('出力不一致','wrong_output'),('IF不適合','interface_mismatch'),('量子ビット数','qubit_count_mismatch'),('実行時（未定義変数）','unknown_error')]
THESIS_TASKS=[('Bell','T1_Bell'),('GHZ','T2_GHZ'),('DJ constant','T3a_DJ_constant'),('DJ balanced','T3b_DJ_balanced'),('BV','T4_BV_011'),('Grover','T5_Grover_11'),('QFT','T6_QFT_001'),('IQFT','T7_IQFT_001'),('QPE','T8_QPE_001'),('Ansatz','T9_Ansatz_001')]

def legacy_thesis():
    d=setup();
    for text,size in [('卒　業　論　文',22),('',10), (TITLE,20),('',10),('2023TC011　伊藤 悠真',14),('指導教員　横山 哲郎',14),('',10),('提出年月：要確認',12),('南山大学 理工学部電子情報工学科',14),('草稿｜2026-10-16版固定集計に基づく',10)]:
        p=para(d,text,size,center=True);p.paragraph_format.space_after=Pt(18)
    d.add_page_break()
    for text,size in [('Graduation Thesis',22),('',10),(TITLE_EN,19),('',10),('2023TC011　伊藤 悠真',14),('Supervisor: 横山 哲郎',14),('Submission month and English name forms: to be confirmed',10),('',10),('Department of Electronics and Communication Technology',12),('Faculty of Science and Technology',12),('Nanzan University',14)]:
        p=para(d,text,size,center=True);p.paragraph_format.space_after=Pt(14)
    sec=d.add_section(WD_SECTION_START.NEW_PAGE);page_numbers(sec,'lowerRoman')
    heading(d,'要約',1)
    for text in JA_SUMMARY.splitlines():para(d,text)
    d.add_page_break();heading(d,'Abstract',1)
    for text in EN_SUMMARY.splitlines():para(d,text)
    d.add_page_break();heading(d,'目次',1)
    p=d.add_paragraph();p.paragraph_format.line_spacing=Pt(13)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(158),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    toc_file=HERE/'toc_cache.json'
    toc=json.loads(toc_file.read_text()) if toc_file.exists() else []
    value='\n'.join(f"{entry['title']}\t{entry['page']}" for entry in toc) or '目次のページ番号はPDFビルド後に更新する．'
    field(p,'TOC \\o "1-2" \\h \\z \\u',value)
    for r in p.runs:setfont(r,10)
    sec=d.add_section(WD_SECTION_START.NEW_PAGE);page_numbers(sec)
    sec.header.is_linked_to_previous=False
    p=sec.header.paragraphs[0];setfont(p.add_run('卒業論文草稿　｜　2026-10-16固定集計'),9)
    for n,(title,sections) in enumerate(CHAPTERS,1):
        hp=heading(d,f'第{n}章　{title}',1)
        if n>1:hp.paragraph_format.page_break_before=True
        for j,(subtitle,paragraphs) in enumerate(sections,1):
            heading(d,f'{n}.{j}　{subtitle}',2)
            for paragraph_index,text in enumerate(paragraphs):
                para(d,text)
                if (n,j,paragraph_index)==(6,1,0):
                    from six_conditions import table_rows
                    table(d,'表6.0　6条件の最終L3成功率（全1,000初期組）',
                          ['条件','L3成功数','分母','最終L3成功率'],table_rows(),widths=[20,42,42,56])
                if (n,j,paragraph_index)==(7,2,0):thesis_tables(d,7)
            # Insert each table near its first textual reference.
            if (n,j)==(3,1):thesis_tables(d,3)
            if (n,j)==(5,1):thesis_tables(d,5)
            if n==6:
                # Tables collected after result subsections are ordered and explicitly referenced.
                pass
        if n==6:
            thesis_tables(d,6)
            # Place tables after the prose that first discusses them.
            body=d._element.body
            for number,target in [(1,'6.2'),(2,'6.3'),(3,'6.4'),(4,'6.5'),(5,'6.6'),(6,'6.7'),(7,'6.7')]:
                caption=next(e for e in body if e.tag==qn('w:p') and ''.join(e.itertext()).startswith(f'表6.{number}　'))
                tbl=caption.getnext(); spacer=tbl.getnext()
                anchor=next(p._p for p in d.paragraphs if p.text.startswith(target+'　'))
                for el in [caption,tbl,spacer]:anchor.addprevious(el)
    d.add_page_break();heading(d,'参考文献',1)
    for i,r in enumerate(REFS,1):para(d,f'[{i}] {r}',indent=False)
    para(d,'文献[1]，[2]，[4]〜[7]の書誌情報と本文で参照した主張の範囲は，原論文の公開ページで確認した．引用番号は本文の初出順である．文献[3]は本研究の一次資料で，集計時点・保存設定・検証記録を固定している．')
    d.add_page_break();heading(d,'付録A　出典と再現性',1)
    manifest=json.loads((ANALYSIS/'manifest.json').read_text())
    for i,s in enumerate(manifest['source_settings'],1):
        experiment_label = ('5モデル追加実験' if 'shared_five_models' in s['source_dir'] else '3モデル本実験' if 'shared_three_models' in s['source_dir'] else '2モデル補完実験')
        heading(d,f'A.{i}　{experiment_label}の保存設定',2)
        para(d,s['source_dir']+'/experiment_config.json',10,indent=False)
        para(d,'設定SHA-256：'+s['config_sha256'],9,indent=False)
        para(d,f"保存schema：{s['schema_version']}．ハーネスcommit：{s['git_commit']}．モデルdigestとプロンプト関数hashは元manifestに保持．",10,indent=False)
    heading(d,'A.4　固定集計と公開範囲',2)
    for text in ['固定集計時点（UTC）：'+SNAP['captured_at_utc'],
       '主出典：docs/presentation_sources/20261016T_reachable_rescues/results_snapshot.json',
       'B4追加実験：independent_4_full_20261020T185811381595Z．固定出典：docs/analysis/independent_4_completed_20261022/．runner_experiment_config.json，runner_manifest.json，summary.csv，pairwise.csv，sanity_report.json，source_manifest.jsonを参照．',
       '全分母：docs/analysis/reachable_rescues_20261016/core/repair_summary.csv',
       '追加消費：同ディレクトリのcost_by_cap_and_stratum.csvとthird_round_summary.csv',
       '片側救済全履歴：discordant_cases.csvとdiscordant_transitions.csv',
       '検証：core/sanity_report.json，sanity_report.json，validation_report.json',
       '介入カタログ：intervention_catalog.csv．事例の記録キー・判定・コードhash・短い抜粋：case_studies.json．固定プロンプト：experiment_prompts.json．追加資料の出典とhash：enrichment_manifest.json．',
       '本稿では検証済み事例の短いコード抜粋と固定プロンプトを掲載する．API応答全文・生成コード全文・生例外文・raw.jsonl・秘密情報は掲載しない．']:
        para(d,text,10,indent=False)
    d.add_page_break();heading(d,'付録B　指標の読み方',1)
    for title,text in [
      ('集計単位','初期組はモデル×タスク×seedの1,000組，B/R/F/D/Pの条件記録は5,000記録，到達roundは9,249 roundである．B4の追加実験1,000組は別に集計する．単位を混ぜて成功率の分母にしない．'),
      ('退行0の意味','R/F/Dは初期成功を保持し，修正中も成功時に停止する．退行0はこの設計を反映しており，自由な修正を続けても正しさが保たれる保証ではない．'),
      ('静的変更の意味','import/APIやゲート名の変更は観察された文字列上の事実であり，その変更単独の因果的効果や意味的正しさを認定したものではない．'),
      ('コストの意味','共有初期を除く入力＋出力tokensと生成要求時間を集計した．再構成経路の計算値で，評価時間・全工程時間・金銭費用とは異なる．'),
    ]:
        heading(d,title,2);para(d,text)
    from thesis_enrichment import add_prompt_appendix
    add_prompt_appendix(d,heading,para)
    # Suppress blue hyperlink styles in the updated TOC as well.
    for style in ['Hyperlink','FollowedHyperlink']:
        if style in d.styles:d.styles[style].font.color.rgb=RGBColor(0,0,0)
    d.save(HERE/'thesis.docx')

def thesis():
    from template_thesis import build_thesis
    build_thesis()


def slides():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    prs.core_properties.title=TITLE;prs.core_properties.author='伊藤 悠真'
    def txt(s,x,y,w,h,text,size=26,bold=False,color='171717'):
        shape=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True
        tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=tf.margin_bottom=Inches(.02)
        for i,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=PPt(10)
            for r in p.runs:
                r.font.name=GOTHIC;r.font.size=PPt(size);r.font.bold=bold;r.font.color.rgb=PColor.from_string(color)
                r._r.get_or_add_rPr().set('lang','ja-JP')
        return shape
    def rect(s,x,y,w,h,shade='EEEEEE'):
        a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));a.fill.solid();a.fill.fore_color.rgb=PColor.from_string(shade);a.line.fill.background();return a
    def base(title,num,subtitle=None):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=PColor(255,255,255)
        txt(s,.6,.35,12.1,.7,title,30,True);rect(s,.63,1.2,12.0,.02,'333333')
        if subtitle:txt(s,.65,1.4,12,1,subtitle,22)
        txt(s,.65,7.08,11.6,.25,'卒業研究発表会　｜　2026-10-16版固定集計 [3]',11,color='555555')
        txt(s,12.2,7.04,.5,.3,str(num),12,color='555555');return s
    notes=[]
    def note(s,time,cues):
        text=f'目安：{time}\n話す要点（読み上げ原稿ではない）\n'+ '\n'.join('・'+x for x in cues)
        s.notes_slide.notes_text_frame.text=text;notes.append(text)
    s=base('初期失敗コードを，どこまで救済できるか',1)
    txt(s,.8,1.8,11.8,1.6,'量子LLM支援プログラミングにおける\n失敗類型と介入カタログの実証研究',28,True)
    txt(s,.85,4.05,11.5,1.2,'2023TC011　伊藤 悠真\n指導教員　横山 哲郎',24)
    txt(s,.85,5.9,11.6,.7,'実行結果提示と自己デバッグを，同じ初期失敗から比較する',23)
    note(s,'0:00–0:15（15秒）',['研究対象はLLM生成のQiskitコード','救済率と救済する対象の違いを調べる'])
    s=base('動くコードでも，期待した出力とは限らない',2)
    for x,label,desc in [(0.8,'読込み','import・APIの失敗'),(4.8,'回路構築','回路を作れない'),(8.8,'出力検証','期待分布と不一致')]:
        rect(s,x,2.05,3.6,1.6);txt(s,x+.15,2.25,3.3,.5,label,27,True);txt(s,x+.15,2.95,3.3,.5,desc,20)
    txt(s,.85,4.4,11.8,1.3,'修正でどれだけ救済できるか\n同じ失敗を救済するのか',30,True)
    txt(s,.85,6.15,11.6,.5,'背景：Self-Refine [1] ／ Self-Debugging [2]（完全再現ではない）',18)
    note(s,'0:15–0:55（40秒）',['実行可能性と出力の正しさを分ける','研究の問いは救済数・救済対象・追加消費'])
    s=base('同じ初期コードから，3種類の修正を比較',3,'5モデル × 10タスク × seed 0〜19 ＝ 1,000初期組')
    rect(s,.8,3,3.1,1.15);txt(s,1,3.2,2.8,.7,'初期生成 B',29,True)
    for y,label in [(2.5,'R　自己見直し'),(3.65,'F　実行結果提示'),(4.8,'D　診断＋説明・原因分析')]:
        txt(s,4.15,y,.6,.6,'→',30);rect(s,4.95,y,7.3,.8);txt(s,5.1,y+.12,6.95,.6,label,25,True)
    txt(s,.85,6.05,11.8,.7,'最大3修正・成功時停止 ／ 1要求1,024 tokens ／ Pは独立生成',21)
    note(s,'0:55–1:45（50秒）',['Rは外部診断なし，FとDは実行診断あり','Dは説明・原因分析も促す','Pを含む5条件で計5,000条件記録'])
    s=base('救済率の分母は，初期失敗545組',4)
    txt(s,.85,1.8,11.7,1,'1,000初期組 ＝ 初期成功455組 ＋ 初期失敗545組',28,True)
    rect(s,.85,3.1,11.6,1.1);txt(s,1.05,3.33,11.1,.65,'救済率 ＝ 初期失敗から最終成功となった組数 ／ 545',26,True)
    txt(s,.85,4.75,11.6,1.5,'成功：読込み・構築・指定入力の期待出力・構造要件\n全入力での正しさを保証する評価ではない',24)
    txt(s,.85,6.35,11.6,.45,'検証済み保存prefix＋必要な部分追試suffixによる現行経路',18)
    note(s,'1:45–2:25（40秒）',['L3を成功基準とする','初期成功は修正せず保持','全件新生成ではなく，到達可能な経路を検証した集計'])
    s=base('FとDは初期失敗の約3分の1を救済',5,'共通の分母：初期失敗545組')
    for y,label,n,rate,col in [(2.6,'R　自己見直し',96,17.6,'AAAAAA'),(3.65,'F　実行結果提示',183,33.6,'333333'),(4.7,'D　自己デバッグ',180,33.0,'777777')]:
        txt(s,.85,y,3.4,.65,label,23);rect(s,4.4,y+.08,rate/40*5.5,.48,col);txt(s,9.3,y,3.1,.75,f'{n}組  {rate:.1f}%',27,True)
    txt(s,.85,6.25,11.7,.5,'今回の対象での観測．モデル・介入の一般的優劣は断定しない．',19)
    note(s,'2:25–3:20（55秒）',['R96/545，F183/545，D180/545','FとDの差は3組','数が近くても対象が同じかは別に調べる'])
    s=base('救済数が近くても，救済する対象は異なる',6)
    for x,w,n,label,col in [(0.85,3.5,'43','Fのみ','EEEEEE'),(4.55,3.5,'140','F/D共通','DDDDDD'),(8.25,3.5,'40','Dのみ','EEEEEE')]:
        rect(s,x,2.05,w,2.2,col);txt(s,x+.2,2.35,w-.4,.9,n+'組',40,True);txt(s,x+.2,3.42,w-.4,.6,label,25)
    txt(s,.85,4.75,11.6,.8,'片側救済83組の全履歴を確認',30,True)
    txt(s,.85,5.6,11.6,1.0,'両条件166履歴・398修正遷移 ／ 双方未救済322組\n複数箇所が同時に変わるため，成功原因は分離できない',21)
    note(s,'3:20–4:10（50秒）',['Fのみ43とDのみ40で計83組','全例を静的に確認した','和集合223組を実行前の選択成功率とは解釈しない'])
    s=base('3回目にも救済がある．追加消費と併せて考える',7)
    txt(s,.85,1.8,5.6,.6,'追加tokens／救済',26,True);txt(s,7.2,1.8,5.4,.6,'3回目の追加救済',26,True)
    for y,label,value in [(2.8,'R','18,852.6'),(3.65,'F','9,333.8'),(4.5,'D','10,976.6')]:
        txt(s,1.05,y,.7,.6,label,26,True);txt(s,2,y,4.6,.6,value,30,True)
    txt(s,7.3,2.8,5.1,1.8,'F　28組\nD　33組',32,True)
    txt(s,7.3,4.65,5.1,1.4,'2回で止めると\nF15.3%・D18.3%の救済を失う',21)
    txt(s,.85,6.35,11.8,.45,'共有初期を除き未救済消費を含む計算値．同じ最大3修正の履歴を区切った比較．',17)
    note(s,'4:10–5:05（55秒）',['消費には失敗した修正も含む','再構成経路の計算値で金銭費用ではない','上限ごとに独立した実験ではなく，一般的な最適回数は未確定'])
    s=base('救済率に加え，「何を救済するか」が重要',8)
    for y,n,text in [(1.85,'1','F183組・D180組を救済（初期失敗545組）'),(3.05,'2','片側救済83組：介入で救済対象が異なる'),(4.25,'3','追加消費と停止回数の両面から検討する')]:
        txt(s,.85,y,.55,.7,n,30,True);txt(s,1.6,y,10.9,.9,text,27,True)
    txt(s,.85,5.8,11.7,.9,'今後：選択・停止方策を事前固定し，別seed・タスクで検証',24)
    note(s,'5:05–6:00（55秒）',['今回の対象・指定入力での結論に限定する','因果効果・統計的有意差は主張しない','最後に今後の検証方針を述べて終了'])
    s=base('補足A　モデル・初期失敗類型で傾向が変わる',9)
    lines=['モデル　失敗分母　R　F　D']+['　'.join(map(str,r)) for r in model_rows()]
    txt(s,.85,1.7,5.8,4.3,'\n'.join(lines),23)
    txt(s,7,1.7,5.5,4.8,'初期import失敗（271組）\nF 58組 ／ D 72組\n\n初期回路構築失敗（126組）\nF 76組 ／ D 68組',23)
    txt(s,.85,6.3,11.6,.5,'Q4/G4/Q9/C7/G3の正式タグは要旨・論文本体のモデル表を参照．',18)
    note(s,'質疑用（本編6分には含めない）',['初期失敗の分母がモデルごとに異なる','類型とモデル・タスクの関連がある'])
    s=base('補足B　評価と経路の範囲',10)
    txt(s,.85,1.75,11.6,4.6,'L3：指定入力の出力＋所定の構造要件\n分布：8,192 shots・全変動距離0.05未満\n検証済み保存prefix＋部分追試293ケースのsuffix\n旧120件は合算しない／Pは独立生成\n共有初期・タスク・seed依存，診断文・負荷差が残る',24)
    txt(s,.85,6.15,11.6,.6,'P：改善124組・悪化106組，最終成功473/1,000組',23,True)
    note(s,'質疑用',['評価は全入力での正しさを保証しない','退行0は成功時停止の設計による','部分追試後に到達しない旧suffixは除外'])
    s=base('補足C　参考文献・固定出典',11)
    for y,text in [(1.7,'[1] Madaan et al., Self-Refine, arXiv:2303.17651v2 (2023).'),(2.8,'[2] Chen et al., Teaching Large Language Models to Self-Debug,\n      arXiv:2304.05128v2 (2023).'),(4.25,'[3] 現行救済経路の固定集計（2026-10-16版）\n      docs/presentation_sources/20261016T_reachable_rescues/'),(5.7,'固定snapshot・全分母CSV・sanity/validation reportを使用．\n本資料の作成に伴う追加実験は行っていない．')]:txt(s,.85,y,11.7,1.1,text,21)
    note(s,'質疑用',['書誌情報と固定出典の案内','提出年月・英語表記・類似レポートは別途確認'])
    prs.save(HERE/'slides.pptx')
    (HERE/'speaker_notes.txt').write_text('6分発表用の要点メモ（原稿の読み上げには使わない）\n本編8枚＝360秒．補足3枚は質疑用．\n\n'+'\n\n'.join(f'{i}枚目\n{x}' for i,x in enumerate(notes,1))+'\n\n練習：スライドだけを見て説明し，6分以内に収める．ゼミの練習で修正する．不明な質問は不明と答え，確認方法を述べる．\n')

def source_record():
    sources=list(B4_SOURCE.glob('*.json'))+list(B4_SOURCE.glob('*.csv'))+[SOURCE/'results_snapshot.json',SOURCE/'validation_report.json',ANALYSIS/'manifest.json',ANALYSIS/'core/repair_summary.csv',ANALYSIS/'core/sanity_report.json',ANALYSIS/'sanity_report.json',ANALYSIS/'core/fd_overlap_summary.csv',ROOT/'harness/evaluator.py',ROOT/'harness/readout.py',ROOT/'tasks/__init__.py']
    from six_conditions import SOURCES as SIX_SOURCES
    sources += SIX_SOURCES
    sources += [HERE/n for n in ['intervention_catalog.csv','case_studies.json','experiment_prompts.json','enrichment_manifest.json']]
    data={'created_at_utc':datetime.now(timezone.utc).isoformat(),'status':'draft','source_snapshot_at_utc':SNAP['captured_at_utc'],
      'source_hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},
      'B4_experiment_id':'independent_4_full_20261020T185811381595Z','B4_snapshot':'docs/analysis/independent_4_completed_20261022','new_experiments':0,'new_model_requests':0,'generated_code_executions':0,
      'literature_checked':[r['url'] for r in json.loads((HERE/'bibliography_sources.json').read_text())['references'] if 'url' in r],
      'thesis_expanded_on':'2026-10-28','thesis_expansion_scope':'固定集計を維持した本文の大幅拡充。新規生成・生成コード実行なし。',
      'pending':['提出年月','英語の著者・教員・所属の正式表記','指導教員による内容と書式の確認','剽窃チェック・類似レポート（審査用と最終版）','Microsoft Word・PowerPoint上での最終レイアウト確認・6分リハーサル']}
    (HERE/'source_manifest.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')

if __name__=='__main__':
    assert SNAP['primary']['initial_groups']==1000 and SNAP['primary']['condition_records']==5000
    assert json.loads((SOURCE/'validation_report.json').read_text())['passed']
    import sys
    if '--thesis-only' in sys.argv:
        thesis()
    else:
        abstract();thesis();slides();source_record()
    print('Created thesis.docx' if '--thesis-only' in sys.argv else 'Created abstract.docx, thesis.docx, slides.pptx, speaker_notes.txt, source_manifest.json')
