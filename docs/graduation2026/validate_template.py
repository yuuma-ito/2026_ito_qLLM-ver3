"""Validate the transcribed-template thesis without claiming original fidelity."""
from datetime import datetime, timezone
import hashlib
import json
import re

from docx import Document
from docx.oxml.ns import qn
import pymupdf
import template_thesis as template

HERE = template.HERE


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate():
    d = Document(HERE/'thesis.docx')
    pdf = pymupdf.open(HERE/'thesis.pdf')
    checks = {}
    near = lambda a,b: abs(a-b)<.08
    checks['a4_portrait_and_requested_margins'] = all(
        all(near(actual,expected) for actual,expected in [(s.page_width.mm,210),(s.page_height.mm,297),
        (s.top_margin.mm,30),(s.bottom_margin.mm,30),(s.left_margin.mm,20),(s.right_margin.mm,20),
        (s.header_distance.mm,0),(s.footer_distance.mm,17.5)]) for s in d.sections)
    checks['single_column'] = all(s._sectPr.find(qn('w:cols')).get(qn('w:num'),'1')=='1' for s in d.sections)
    checks['empty_header'] = all(not any(p.text.strip() for p in s.header.paragraphs) for s in d.sections)
    checks['page_field_right_aligned_including_cover'] = all(
        not s.different_first_page_header_footer and s.footer.paragraphs[0].alignment==2 and
        any(e.text.strip()=='PAGE' for e in s.footer._element.iter(qn('w:instrText'))) for s in d.sections)
    for style,size,jp,latin in [('Normal',10.5,'MS 明朝','Century'),('Heading 1',18,'MS Pゴシック','Arial'),('Heading 2',14,'MS Pゴシック','Arial'),('Heading 3',12,'MS Pゴシック','Arial')]:
        st=d.styles[style];rf=st.element.rPr.rFonts
        checks[style+'_font'] = near(st.font.size.pt,size) and rf.get(qn('w:eastAsia'))==jp and rf.get(qn('w:ascii'))==latin
    checks['normal_spacing_not_fixed'] = d.styles['Normal'].element.pPr.find(qn('w:spacing')) is None
    checks['heading_multilevel_numbering'] = all(d.styles[f'Heading {i}'].element.pPr.find(qn('w:numPr')).find(qn('w:numId')).get(qn('w:val'))=='90' for i in range(1,4))
    definitions=next(e for e in d.part.numbering_part.element if e.tag==qn('w:abstractNum') and e.get(qn('w:abstractNumId'))=='90')
    checks['number_patterns'] = [e.find(qn('w:lvlText')).get(qn('w:val')) for e in definitions if e.tag==qn('w:lvl')]==['%1','%1.%2','%1.%2.%3']
    headings=[p for p in d.paragraphs if p.style.name.startswith('Heading ')]
    checks['heading_numbers_not_typed'] = all(not re.match(r'^(?:第\d|\d+\.|付録[A-Z])',p.text) for p in headings)
    checks['chapter_order'] = [p.text for p in headings if p.style.name=='Heading 1']==['はじめに','関連研究','評価・実験方法','結果','考察','おわりに','参考文献','付録']
    checks['heading_keep_with_next'] = all(d.styles[f'Heading {i}'].paragraph_format.keep_with_next for i in range(1,4))
    checks['chapter_page_breaks'] = d.styles['Heading 1'].paragraph_format.page_break_before
    instructions=[e.text.strip() for e in d.element.iter(qn('w:instrText'))]
    checks['editable_toc_levels_1_to_3'] = any('TOC' in i and '"1-3"' in i for i in instructions)
    cache=json.loads((HERE/'toc_cache.json').read_text())
    checks['all_toc_entries_match_pdf']=len(cache)==len(headings) and cache==template.toc_entries()
    checks['toc_dot_leaders_and_right_tabs']=all(
        p.paragraph_format.tab_stops[0].alignment==2 and p.paragraph_format.tab_stops[0].leader==1
        for p in d.paragraphs if '\t' in p.text)
    checks['sequential_table_fields']=sum(i.startswith('SEQ Table ') for i in instructions)==len(d.tables)==11
    checks['editable_table_references']=any(i.startswith('REF Table') for i in instructions)
    checks['four_figures_with_sequential_fields'] = len(d.inline_shapes) == sum(i.startswith('SEQ Figure ') for i in instructions) == 4
    checks['editable_figure_references'] = all(any(i.startswith('REF Figure'+str(n)+' ') for i in instructions) for n in [1,2,3,4])
    figure_caps=[p for p in d.paragraphs if any('SEQ Figure' in e.text for e in p._p.iter(qn('w:instrText')))]
    checks['figure_captions_below_images'] = len(figure_caps)==4 and all(p._p.getprevious() is not None and bool(p._p.getprevious().xpath('.//w:drawing')) for p in figure_caps)
    checks['appendix_titles_distinct'] = [h['title'] for h in cache if h['number'] in ['8.1.1','8.1.2','8.1.3']] == ['5モデル追加実験の保存設定','3モデル本実験の保存設定','2モデル補完実験の保存設定']
    checks['table_widths_within_170mm']=all(sum(c.width.mm for c in t.columns)<=170.1 for t in d.tables)
    checks['table_headers_repeat']=all(t.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None for t in d.tables)
    caption_ps=[p for p in d.paragraphs if any('SEQ Table' in e.text for e in p._p.iter(qn('w:instrText')))]
    checks['captions_before_table_and_kept_with_next']=all(p._p.getnext().tag==qn('w:tbl') and p.paragraph_format.keep_with_next for p in caption_ps)
    tables=[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    from six_conditions import table_rows, DEFINITIONS
    checks['six_condition_table_matches_csv'] = tables[2][1:] == table_rows()
    checks['six_condition_definitions'] = tables[0][1:] == DEFINITIONS
    overall = {r['condition']:r for r in template.content.b4_rows('overall')}
    expected = [[c,overall[c]['initial_failures'],overall[c]['successes_from_initial_failure'],f"{100*float(overall[c]['success_from_initial_failure_rate']):.1f}%"] for c in ['R','F','D','B4']]
    checks['overall_table_matches_frozen_csv'] = tables[3][1:] == expected and tables[3][0] == ['条件','初期失敗数','追加成功数','追加成功率']
    for index,grouping,labels in [(4,'model',[(short,'ollama:'+model) for short,model in template.content.MODELS]),(5,'initial_failure',template.content.THESIS_ERRORS),(6,'task',template.content.THESIS_TASKS)]:
        expected=[]
        for label,value in labels:
            rr={r['condition']:r for r in template.content.b4_rows(grouping,value)}
            expected.append([label,rr['B4']['initial_failures']]+[rr[c]['successes_from_initial_failure'] for c in ['R','F','D','B4']])
        checks[grouping+'_table_matches_frozen_csv'] = tables[index][1:] == expected
    checks['overlap_table_matches_counts'] = tables[7][1:] == [['140','43','40','322']]
    checks['additional_cost_table_matches_csv'] = tables[8][1:] == [[c,overall[c]['additional_requests'],f"{int(overall[c]['additional_tokens']):,}",f"{float(overall[c]['tokens_per_added_success']):,.1f}"] for c in ['R','F','D','B4']]
    third={r['condition_id']:r for r in template.content.SNAP['third_round']}
    checks['third_round_table_matches_snapshot'] = tables[9][1:] == [[c,f"{third[k]['additional_rescues']}/{third[k]['at_risk_after_2']}",f"{third[k]['third_round_tokens']:,}",f"{third[k]['third_round_seconds']:,.1f}",f"{third[k]['loss_if_stop_at_2']*100:.1f}%"] for c,k in template.content.CONDITIONS]
    from thesis_enrichment import catalog_rows
    checks['catalog_table_matches_frozen_counts'] = tables[10][1:] == catalog_rows()
    checks['tables_not_justified'] = all(p.alignment in [0,1,2] for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs)
    checks['body_left_aligned'] = all(p.alignment == 0 for p in d.paragraphs if p.text.strip() and p.style.name == 'Normal' and not any('SEQ Table' in e.text or 'SEQ Figure' in e.text for e in p._p.iter(qn('w:instrText'))) and '\t' not in p.text and not any(t in p.text.replace('\n','') for t in ['卒　業　論　文',template.content.TITLE,'2023TC011','指導教員','2027年1月','南山大学 理工学部']))
    checks['B4_pairwise_counts_match_prose'] = {r['comparator']:[int(r.get(k) or 0) for k in ['both','B4_only',r['comparator']+'_only','neither']] for r in template.content.csv.DictReader((template.content.B4_SOURCE/'pairwise.csv').open())} == {'R':[69,148,27,301],'F':[131,86,52,276],'D':[120,97,60,268]}
    checks['B4_sanity_passed'] = json.loads((template.content.B4_SOURCE/'sanity_report.json').read_text())['passed']
    norm=lambda text:''.join(text.split())
    texts=[p.get_text() for p in pdf]
    checks['cover_identity_date']=all(norm(value) in norm(texts[0]) for value in ['卒業論文',template.content.TITLE,'2023TC011 伊藤 悠真','横山 哲郎','2027年1月'])
    checks['summary_second_page']=all(norm(value) in norm(texts[1]) for value in ['南山大学 理工学部 電子情報学科','要約','545','183','180'])
    label_lines=[]
    for page in pdf:
        lines=[''.join(sp['text'] for sp in line['spans']) for b in page.get_text('dict')['blocks'] for line in b.get('lines',[])]
        for pos,line in enumerate(lines):
            if norm(line) == norm('表6 初期失敗類型別の追加成功数'):
                label_lines.extend(lines[pos+1:])
    checks['failure_labels_not_wrapped'] = all(any(norm(label) in norm(line) for line in label_lines) for label,_ in template.content.THESIS_ERRORS)
    checks['toc_third_page']='目次' in texts[2]
    first_body_page = int(cache[0]['page'])
    checks['no_blank_pages'] = not [i for i,text in enumerate(texts,1) if len(norm(re.sub(r'\d+\s*$', '', text))) < 5]
    checks['toc_may_span_multiple_pages'] = first_body_page >= 4 and all(str(entry['page']) in ''.join(texts[2:first_body_page-1]) for entry in cache)
    checks['continuous_visible_decimal_pages']=all(text.rstrip().splitlines()[-1]==str(i) for i,text in enumerate(texts,1))
    checks['pdf_a4']=all(abs(p.rect.width-595.28)<.6 and abs(p.rect.height-841.89)<.6 for p in pdf)
    outside=[];margin=[]
    for i,p in enumerate(pdf,1):
        for b in p.get_text('rawdict')['blocks']:
            for line in b.get('lines',[]):
                for sp in line['spans']:
                    # Justified PDF spans can include trailing spaces beyond the
                    # print area. Inspect visible characters, not space advance.
                    chars = [c for c in sp['chars'] if c['c'].strip()]
                    if not chars:
                        continue
                    r = pymupdf.Rect(chars[0]['bbox'])
                    for c in chars[1:]:
                        r |= pymupdf.Rect(c['bbox'])
                    if r.x0<-.5 or r.y0<-.5 or r.x1>p.rect.width+.5 or r.y1>p.rect.height+.5:outside.append(i)
                    if r.x0<56 or r.x1>539.5:margin.append({'page':i,'text':''.join(c['c'] for c in chars)})
    checks['no_text_outside_page']=not outside
    checks['text_within_left_right_margins']=not margin
    xrefs={f[0] for p in pdf for f in p.get_fonts()}
    checks['pdf_fonts_embedded']=all(bool(pdf.extract_font(x)[3]) for x in xrefs)
    checks['no_replacement_characters']=all('\ufffd' not in t for t in texts)
    checks['no_field_errors']=not any(s in '\n'.join(texts) for s in ['Error!','エラー!','参照元が見つかりません','目次はPDF'])
    checks['under_10mb']=(HERE/'thesis.pdf').stat().st_size<10_000_000
    manifest=json.loads((HERE/'source_manifest.json').read_text())
    checks['frozen_source_hashes_unchanged']=all(digest(template.content.ROOT/p)==sha for p,sha in manifest['source_hashes'].items())
    # Square brackets in literal experiment prompts (e.g. qubit indices) are
    # not bibliography citations. Limit this check to manuscript and references.
    citation_paragraphs=[]
    for p in d.paragraphs:
        if p.style.name=='Heading 1' and p.text=='付録':
            break
        citation_paragraphs.append(p.text)
    alltext='\n'.join(citation_paragraphs)
    refs=set(re.findall(r'\[(\d+)\]',alltext))
    checks['reference_numbers_consistent'] = refs == {str(i) for i in range(1,len(template.content.REFS)+1)}
    def subsection(title):
        result=[];active=False
        for p in d.paragraphs:
            if p.style.name.startswith('Heading'):
                if active:break
                active=p.text==title
            elif active:result.append(p.text)
        return '\n'.join(result)
    questions=subsection('研究目的と研究課題')
    comparison=subsection('修正と独立再生成の比較から得られる知見')
    resources=subsection('資源制約と停止方策への含意')
    checks['RQ1_to_RQ3_explicit_in_introduction'] = all(f'RQ{i}（' in questions for i in [1,2,3])
    checks['RQ_answers_in_discussion'] = all(f'RQ{i}への回答' in comparison for i in [1,2]) and 'RQ3への回答' in resources
    checks['RQ3_includes_B4_cost_from_frozen_csv'] = all(t in resources for t in [
        f"B4の追加tokens／追加成功は{float(overall['B4']['tokens_per_added_success']):,.1f}",
        f"総追加{int(overall['B4']['additional_tokens']):,} tokens",
        f"追加成功{overall['B4']['successes_from_initial_failure']}組で割った値"])
    checks['order_of_metrics_explained'] = 'まず全対象の最終的な到達状況を比較し，次に初期成功の保持を除いて' in alltext
    conclusion=subsection('研究成果')
    checks['conclusion_states_catalog_and_no_policy_validation'] = all(t in conclusion for t in [
        '主要な初期失敗3類型（import失敗，回路構築失敗，出力不一致）',
        '初期失敗数とF/Dによる修正の救済率，B4による独立再生成の追加成功率',
        '自動選択方策の有効性を実証したものではない'])
    limitations=subsection('妥当性と限界')
    checks['critical_limitations_retained'] = all(t in limitations for t in [
        'L3は指定入力の出力と構造要件への適合', '実行時期などの交絡は残る',
        '量子変換の意味的差を測定していない', '複数同時変更の寄与も分離できない']) and all(t in comparison for t in [
        '投入情報，tokens，seed規則，実行時期は一致しない', '成功率の差を介入の因果効果と解釈しない'])
    visual_file=HERE/'thesis_visual_review.json'
    visual=json.loads(visual_file.read_text()) if visual_file.exists() else {'status':'pending'}
    visual_current=visual.get('pdf_sha256')==digest(HERE/'thesis.pdf') and visual.get('pages_reviewed')==list(range(1,len(pdf)+1))
    result={
        'checked_at_utc':datetime.now(timezone.utc).isoformat(), 'scope':'thesis.docx and thesis.pdf only; separate abstract and slides retained',
        'automated_checks_passed':all(checks.values()),'checks':checks,
        'documents':{'thesis':{'pages':len(pdf),'pdf_bytes':(HERE/'thesis.pdf').stat().st_size,'pdf_sha256':digest(HERE/'thesis.pdf')}},
        'visual_review':visual if visual_current else {'status':'pending_or_stale'},
        'original_template_comparison':'not performed; original .doc not accessible',
        'font_fidelity':'DOCX specifies requested fonts; LibreOffice PDF uses substitutes. Final Microsoft Word rendering unverified.',
        'format_spec':'format_spec.json', 'source_snapshot_at_utc':manifest['source_snapshot_at_utc'],
        'new_experiments':0, 'new_model_requests':0,'generated_code_executions':0,
        'artifact_sha256':{name:digest(HERE/name) for name in ['thesis.docx','thesis.pdf','build.py','template_thesis.py','thesis_expansion.py','thesis_language.py','thesis_figures.py','validate_figures.py','figure_validation_report_20261028.json','figures/figure_manifest.json','convert_thesis.py','validate_template.py','validate_expansion.py','thesis_headings.json','toc_cache.json','format_spec.json','source_manifest.json','bibliography_sources.json','expansion_report_20261028.json','revision_report_20261028.json','work_notes.txt','thesis_visual_review.json','README.md','build_research_evidence.py','thesis_enrichment.py','validate_enrichment.py','intervention_catalog.csv','case_studies.json','experiment_prompts.json','enrichment_manifest.json','enrichment_validation_report.json','MANUSCRIPT_STATE.md','skill_edit_report_20261029.json','prompt_layout_revision_20261029.json','six_conditions.py','six_conditions_validation.json','validate_six_conditions.py','rq_revision_validation.json']},
        'out_of_margin_spans':margin,
        'pending':['Original template comparison, especially cover coordinates and inherited paragraph styles','Formal approval of thesis title','Final rendering and updated fields in Microsoft Word with the specified fonts','Supervisor review of content and references']}
    (HERE/'validation_report.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':all(checks.values()),'failed':[k for k,v in checks.items() if not v],'pages':len(pdf),'visual_review_current':visual_current},ensure_ascii=False))
    if not all(checks.values()):raise SystemExit(1)


if __name__=='__main__':validate()
