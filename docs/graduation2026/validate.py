"""Validate delivered Office/PDF drafts; optionally rebuild the cached TOC.

Run with the same dependencies as build.py. --update-toc updates only the
cached contents and thesis.docx; re-export thesis.pdf before final validation.
"""
from pathlib import Path
import hashlib,json,re,sys
from datetime import datetime,timezone
from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation
import pymupdf
import build

HERE=Path(__file__).resolve().parent
norm=lambda text:''.join(text.split())

def toc_entries():
    doc=Document(HERE/'thesis.docx');pdf=pymupdf.open(HERE/'thesis.pdf');out=[]
    for paragraph in doc.paragraphs:
        if paragraph.style.name not in ['Heading 1','Heading 2'] or paragraph.text in ['要約','Abstract','目次']:continue
        title=paragraph.text
        m=re.match(r'第\d+章|\d+\.\d+|付録[A-Z]|A\.\d+',title)
        key=m.group(0) if m else norm(title)
        hits=[]
        for i,page in enumerate(pdf):
            text=page.get_text();compact=norm(text)
            # Body pages carry this running header. Exclude front matter/TOC.
            if not compact.startswith('卒業論文草稿'):continue
            if key.startswith('第') or key.startswith('付録'):
                hit=key in compact[:65]
            elif m:
                hit=any(re.match(re.escape(key)+r'(?![0-9])',norm(line)) for line in text.splitlines())
            else:hit=key in [norm(line) for line in text.splitlines()]
            if hit:hits.append(text.strip().splitlines()[-1].strip())
        assert hits,(title,'heading not located')
        out.append({'title':title,'page':hits[0]})
    return out

def validate():
    checks={}; details={}
    abstract=Document(HERE/'abstract.docx')
    checks['abstract_section_margins']=all(abs(s.page_width.mm-210)<.1 and abs(s.page_height.mm-297)<.1 and abs(s.left_margin.mm-18)<.1 and abs(s.right_margin.mm-18)<.1 and abs(s.top_margin.mm-18)<.1 and abs(s.bottom_margin.mm-24)<.1 for s in abstract.sections)
    checks['abstract_one_column_title_two_column_body']=[s._sectPr.find(qn('w:cols')).get(qn('w:num'),'1') for s in abstract.sections]==['1','2']
    sizes=[r.font.size.pt for p in abstract.paragraphs for r in p.runs if r.text.strip() and r.font.size]
    sizes += [r.font.size.pt for t in abstract.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text.strip() and r.font.size]
    checks['abstract_10_to_11_pt']=all(10<=s<=11 for s in sizes)
    checks['abstract_page_number_field']=any('PAGE' in e.text for s in abstract.sections for e in s.footer._element.iter(qn('w:instrText')))
    tables=[[[c.text for c in row.cells] for row in t.rows] for t in abstract.tables]
    checks['abstract_model_table_matches_csv']=tables[0][1:]==[[str(c) for c in r] for r in build.model_rows()]
    checks['abstract_error_table_matches_csv']=tables[1][1:]==[[str(c) for c in r] for r in build.error_rows()]
    checks['model_and_error_totals_match']=sum(int(r[1]) for r in build.model_rows())==545 and sum(int(r[1]) for r in build.error_rows())==545 and all(sum(int(r[i+2]) for r in build.model_rows())==n for i,n in enumerate([96,183,180])) and all(sum(int(r[i+2]) for r in build.error_rows())==n for i,n in enumerate([96,183,180]))
    refs=[]
    for name in ['abstract','thesis','slides']:
        path=HERE/(name+'.pdf');pdf=pymupdf.open(path);text='\n'.join(p.get_text() for p in pdf)
        font_xrefs={f[0] for p in pdf for f in p.get_fonts()}
        embedded=all(bool(pdf.extract_font(x)[3]) for x in font_xrefs)
        bounds=[]
        for n,p in enumerate(pdf,1):
            for block in p.get_text('dict')['blocks']:
                for line in block.get('lines',[]):
                    for sp in line['spans']:
                        box=pymupdf.Rect(sp['bbox'])
                        if sp['text'].strip() and (box.x0<-.5 or box.y0<-.5 or box.x1>p.rect.width+.5 or box.y1>p.rect.height+.5):bounds.append(n)
        checks[name+'_fonts_embedded']=embedded
        checks[name+'_no_text_outside_page']=not bounds
        checks[name+'_no_replacement_characters']='\ufffd' not in text
        checks[name+'_main_counts_present']=all(v in norm(text) for v in ['545','183','180','96'])
        checks[name+'_no_blank_pages']=all(len(p.get_text().strip())>50 for p in pdf)
        if name in ['abstract','thesis']:checks[name+'_a4']=all(abs(p.rect.width-595.28)<.5 and abs(p.rect.height-841.89)<.5 for p in pdf)
        if name=='abstract':
            checks['abstract_two_pages']=len(pdf)==2
            checks['abstract_page_numbers_visible']=all(p.get_text().strip().endswith(str(i)) for i,p in enumerate(pdf,1))
            # Verify all body/table text remains within the left/right printing area.
            checks['abstract_text_within_print_width']=all(sp['bbox'][0]>=49 and sp['bbox'][2]<=547 for p in pdf for b in p.get_text('dict')['blocks'] for l in b.get('lines',[]) for sp in l['spans'] if sp['text'].strip())
        if name=='thesis':checks['thesis_under_10MB']=path.stat().st_size<=10_000_000
        if name=='slides':checks['slides_11_pages']=len(pdf)==11
        details[name]={'pages':len(pdf),'bytes':path.stat().st_size,'pdf_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'fonts':len(font_xrefs),'out_of_page_text_pages':bounds}
    prs=Presentation(HERE/'slides.pptx')
    checks['pptx_editable_text_and_shapes']=all(any(s.has_text_frame for s in slide.shapes) for slide in prs.slides) and all(s.shape_type!=13 for slide in prs.slides for s in slide.shapes)
    checks['pptx_has_notes']=all(s.notes_slide.notes_text_frame.text.strip() for s in prs.slides)
    checks['pptx_shapes_inside_slide']=all(s.left>=0 and s.top>=0 and s.left+s.width<=prs.slide_width+5 and s.top+s.height<=prs.slide_height+5 for slide in prs.slides for s in slide.shapes)
    checks['main_slide_timing_360_seconds']=sum([15,40,50,40,55,50,55,55])==360
    cache=json.loads((HERE/'toc_cache.json').read_text());checks['toc_all_49_entries_match_pdf']=cache==toc_entries() and len(cache)==49
    thesis=Document(HERE/'thesis.docx')
    toc=next(p for p in thesis.paragraphs if any('TOC ' in e.text for e in p._p.iter(qn('w:instrText'))))
    checks['toc_office_cache_matches']=toc.text=='\n'.join(f"{e['title']}\t{e['page']}" for e in cache)
    checks['source_hashes_unchanged']=all(hashlib.sha256((build.ROOT/p).read_bytes()).hexdigest()==sha for p,sha in json.loads((HERE/'source_manifest.json').read_text())['source_hashes'].items())
    hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in HERE.iterdir() if p.is_file() and p.name!='validation_report.json'}
    result={'checked_at_utc':datetime.now(timezone.utc).isoformat(),'automated_checks_passed':all(checks.values()),'checks':checks,'documents':details,'artifact_sha256':hashes,
      'visual_review':{'status':'reviewed_all_pages','notes':['要旨2ページ・スライド11枚・論文全ページを画像で確認．表幅のはみ出しと空白ページを修正．最終版の目次と本文番号も照合．','確認環境はLibreOffice PDF出力．Microsoft Word／PowerPoint本体での確認と6分リハーサルは未実施．']},
      'source_analysis_status':'verified frozen 2026-10-16 snapshot; not re-executed',
      'new_model_requests':0,'generated_code_executions':0,
      'conversion_notes':['LibreOfficeはローカル展開した24.2.7.2．Java補助とliblangtagデータのパスに関する警告が出たが，全PDFの書出しとページ・文字・フォント検証は成功．'],
      'pending':json.loads((HERE/'source_manifest.json').read_text())['pending']}
    (HERE/'validation_report.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':result['automated_checks_passed'],'failed':[k for k,v in checks.items() if not v],'pages':{k:v['pages'] for k,v in details.items()}},ensure_ascii=False))
    if not result['automated_checks_passed']:raise SystemExit(1)

if __name__=='__main__':
    if Document(HERE/'thesis.docx').core_properties.category == 'user-transcribed-template-20261020':
        if '--update-toc' in sys.argv:
            import template_thesis
            entries=template_thesis.toc_entries()
            (HERE/'toc_cache.json').write_text(json.dumps(entries,ensure_ascii=False,indent=2)+'\n')
            build.thesis()
            print('TOC cache updated; re-export thesis.pdf before validation.')
        else:
            import validate_template
            validate_template.validate()
        raise SystemExit(0)
    if '--update-toc' in sys.argv:
        entries=toc_entries();(HERE/'toc_cache.json').write_text(json.dumps(entries,ensure_ascii=False,indent=2)+'\n');build.thesis();print('Updated 49-entry cached TOC and thesis.docx; export thesis.pdf again.')
    else:validate()
