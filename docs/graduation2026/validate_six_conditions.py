"""Check six-condition reporting, preserved analyses and distinct denominators."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
import pymupdf
import six_conditions as six

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]


def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def norm(text):
    return re.sub(r'\s','',text)


def validate():
    values=six.values()
    doc=Document(HERE/'thesis.docx');pdf=pymupdf.open(HERE/'thesis.pdf')
    text=norm(''.join(p.get_text() for p in pdf))
    tables=[[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    manifest=json.loads((HERE/'figures/figure_manifest.json').read_text())
    expected=[455,473,551,638,635,672]
    instructions=[e.text.strip() for e in doc.element.iter(qn('w:instrText'))]
    captions={kind:[p.text for p in doc.paragraphs if any(e.text.strip().startswith('SEQ '+kind+' ') for e in p._p.iter(qn('w:instrText')))] for kind in ['Table','Figure']}
    refs={kind:{int(re.match(r'REF '+kind+r'(\d+)',i)[1]) for i in instructions if i.startswith('REF '+kind)} for kind in ['Table','Figure']}
    final_pdf=pymupdf.open(HERE/'figures/final_success_rates.pdf')
    final_text=norm(''.join(p.get_text() for p in final_pdf))
    added_pdf=pymupdf.open(HERE/'figures/additional_success_rates.pdf')
    added_text=norm(''.join(p.get_text() for p in added_pdf))
    checks={
        'six_counts_match_reference_and_fixed_csv':[r['successes'] for r in values]==expected,
        'six_rates_computed_from_1000':all(r['denominator']==1000 and r['rate_percent']==r['successes']/10 for r in values),
        'six_condition_definitions':tables[0][1:]==six.DEFINITIONS,
        'six_final_success_table':tables[2][1:]==six.table_rows(),
        'table_and_graph_use_identical_values':manifest['final_rates']==values,
        'new_graph_labels_counts_rates_denominators':all(norm(f"{r['successes']}/1,000（{r['rate_percent']:.1f}%）") in final_text for r in values) and all(r['condition'] in final_text for r in values),
        'original_added_success_analysis_retained':[r['added_successes'] for r in manifest['rates']]==[96,183,180,217] and all(r['denominator']==545 for r in manifest['rates']),
        'only_repair_and_B4_in_added_success_graph':[r['condition'] for r in manifest['rates']]==['R','F','D','B4'],
        'distinct_graph_denominators':norm('分母：全1,000初期組') in final_text and norm('分母：初期失敗545組') in added_text,
        'all_11_tables_and_4_figures':len(tables)==11 and len(doc.inline_shapes)==4,
        'captions_numbered_in_order':all([int(re.match(r'(?:表|図)(\d+)',c)[1]) for c in captions[kind]]==list(range(1,total+1)) for kind,total in [('Table',11),('Figure',4)]),
        'all_caption_numbers_referenced':refs=={'Table':set(range(1,12)),'Figure':set(range(1,5))},
        'six_comparison_precedes_added_successes':text.index(norm('表3と図2に'))<text.index(norm('表4に，この同じ失敗集合')),
        'B_P_B4_initial_code_distinction':all(norm(t) in text for t in ['B4もBの初期コードを共有','PはBとは異なる初期コードを独立に生成','初回1回と最大3回の追加生成']),
        'noncausal_interpretation_and_nonidentical_budgets':all(norm(t) in text for t in ['投入情報，tokens，seed規則，実行時期','成功率の差を介入の因果効果と解釈しない']),
        'summary_and_conclusion_include_six_final_rates':True,
    }
    summary=norm(pdf[1].get_text())
    conclusion=norm(''.join(p.get_text() for p in pdf if '研究成果' in p.get_text()))
    for r in values:
        value=norm(f"{r['condition']}{r['rate_percent']:.1f}%")
        checks['summary_and_conclusion_include_six_final_rates'] &= value in summary and value in conclusion
    before=ROOT/'.cache/thesis_six_conditions_20261029/thesis.docx'
    if before.exists():
        old=Document(before)
        old_tables=[[[c.text for c in r.cells] for r in t.rows] for t in old.tables]
        checks['all_prior_numerical_tables_preserved']=old_tables[1:]==tables[1:2]+tables[3:]
        def media_hashes(path):
            with zipfile.ZipFile(path) as archive:
                return {hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/')}
        checks['all_prior_embedded_figures_preserved']=media_hashes(before)<=media_hashes(HERE/'thesis.docx')
    # Read-only hashes: no re-evaluation, generation or aggregate mutation.
    frozen=json.loads((HERE/'source_manifest.json').read_text())['source_hashes']
    checks['frozen_source_hashes_unchanged']=all(digest(ROOT/p)==h for p,h in frozen.items())
    evidence=json.loads((HERE/'enrichment_manifest.json').read_text())['outputs_sha256']
    checks['prompt_case_catalog_evidence_unchanged']=all(digest(HERE/p)==h for p,h in evidence.items())
    report={'checked_at_utc':datetime.now(timezone.utc).isoformat(),'passed':all(checks.values()),'checks':checks,
            'values':values,'fixed_csv_sources':{str(p.relative_to(ROOT)):digest(p) for p in six.SOURCES},
            'captions':captions,'pages':len(pdf),'pdf_sha256':digest(HERE/'thesis.pdf'),'docx_sha256':digest(HERE/'thesis.docx'),
            'new_experiments':0,'re_evaluations':0,'generated_code_executions':0,
            'commit_message_proposal':'docs: 卒論に6条件の最終L3成功率比較を追加'}
    (HERE/'six_conditions_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':report['passed'],'failed':[k for k,v in checks.items() if not v],'pages':len(pdf)},ensure_ascii=False))
    assert report['passed']


if __name__=='__main__':validate()
