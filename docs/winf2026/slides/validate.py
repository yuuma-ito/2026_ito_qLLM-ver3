"""Validate slide artifacts and source hashes; render page previews for visual review.
Run with PYTHONPATH=.cache/graduation_tools .venv/bin/python docs/winf2026/slides/validate.py
"""
from pathlib import Path
import hashlib
import json
import re
from datetime import datetime, timezone

import pymupdf as fitz
from PIL import Image, ImageDraw
from pptx import Presentation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
CACHE = ROOT/'.cache/winf_slides_20261028'
CACHE.mkdir(parents=True,exist_ok=True)
manifest=json.loads((HERE/'source_manifest.json').read_text())
prs=Presentation(HERE/'slides.pptx'); pdf=fitz.open(HERE/'slides.pdf')
hashof=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
normalize=lambda s:re.sub(r'\s+','',s).replace('−','-').replace('–','-').replace('ﬁ','fi').replace('ﬂ','fl')
checks={
    'sources_unchanged':all(hashof(ROOT/p)==h for p,h in manifest['source_sha256'].items()),
    '17_slides_and_pages':len(prs.slides)==len(pdf)==17,
    '13_main_4_appendix':manifest['main_slides']==13 and manifest['appendix_slides']==4,
    'duration_10_to_12_minutes':600<=manifest['duration_seconds']<=720,
    'fixed_title':prs.core_properties.title=='量子LLM支援プログラミングにおける失敗類型と介入カタログの実証研究',
    'speaker_notes_every_slide':all(s.notes_slide.notes_text_frame.text.strip() for s in prs.slides),
}
missing=[];offpage=[];page_info=[]
for i,(slide,page) in enumerate(zip(prs.slides,pdf),1):
    rendered=normalize(page.get_text())
    for shape in slide.shapes:
        texts=[]
        if shape.has_text_frame:texts.append(shape.text)
        if shape.has_table:texts.extend(cell.text for rr in shape.table.rows for cell in rr.cells)
        for txt in texts:
            if normalize(txt) and normalize(txt) not in rendered:missing.append({'slide':i,'text':txt})
        if min(shape.left,shape.top)<0 or shape.left+shape.width>prs.slide_width+100 or shape.top+shape.height>prs.slide_height+100:
            offpage.append({'slide':i,'shape':shape.name})
    bounds=[]
    for b in page.get_text('dict')['blocks']:
        for line in b.get('lines',[]):
            for span in line.get('spans',[]):
                r=fitz.Rect(span['bbox']);bounds.append(list(r))
                if r.x0<0 or r.y0<0 or r.x1>page.rect.width+.5 or r.y1>page.rect.height+.5:offpage.append({'slide':i,'text':span['text']})
    pix=page.get_pixmap(matrix=fitz.Matrix(1.15,1.15));pix.save(CACHE/f'slide_{i:02}.png')
    page_info.append({'slide':i,'width_pt':page.rect.width,'height_pt':page.rect.height,'text_spans':len(bounds)})
checks['all_pptx_text_present_in_pdf']=not missing
checks['all_elements_inside_slide']=not offpage
checks['all_pages_16_9']=all(abs(p.rect.width/p.rect.height-16/9)<.002 for p in pdf)
checks['fonts_embedded']=all(pdf.extract_font(f[0])[3] for p in pdf for f in p.get_fonts() if f[1]!='n/a')
# Reconcile the headline numbers directly with CSV-derived frozen source values.
snap=json.loads((HERE.parent/'results_snapshot_20261028.json').read_text())
checks['headline_counts_match_snapshot']=all(str(snap['overall'][c]['successes_from_initial_failure']) in normalize(pdf[5 if c!='B4' else 8].get_text()) for c in ['R','F','D','B4'])
# Contact sheets: four slides each, for full visual inspection.
for start in range(1,18,4):
    contact=Image.new('RGB',(1600,960),'#DCE5EB');draw=ImageDraw.Draw(contact)
    for k,n in enumerate(range(start,min(start+4,18))):
        im=Image.open(CACHE/f'slide_{n:02}.png').convert('RGB');im.thumbnail((780,439))
        x=10+(k%2)*800;y=25+(k//2)*480
        contact.paste(im,(x,y));draw.text((x,y-18),f'Slide {n:02}',fill='black')
    contact.save(CACHE/f'contact_{start:02}.png')
report={'validated_at_utc':datetime.now(timezone.utc).isoformat(),'automated_checks_passed':all(checks.values()),'checks':checks,'missing_text':missing,'offpage_elements':offpage,'pages':page_info,'artifact_sha256':{name:hashof(HERE/name) for name in ['slides.pptx','slides.pdf','speaker_notes.txt','source_manifest.json','build.py','validate.py']},'visual_review':{'status':'pending','preview_directory':str(CACHE.relative_to(ROOT))},'scope':'資料作成のみ。新規生成・生成コード実行0。'}
(HERE/'validation_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'passed':report['automated_checks_passed'],'checks':checks,'missing_text':missing,'offpage':offpage},ensure_ascii=False,indent=2))
assert report['automated_checks_passed']
