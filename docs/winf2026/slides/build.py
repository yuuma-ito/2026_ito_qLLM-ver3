"""Build editable WiNF slides from frozen, public aggregate sources only.

Usage: PYTHONPATH=.cache/graduation_tools .venv/bin/python docs/winf2026/slides/build.py
       Add --pdf to export with LibreOffice (system or repository-local cache).
"""
from pathlib import Path
import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
import os
import shutil
import subprocess
import tempfile

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SNAP_PATH = HERE.parent / 'results_snapshot_20261028.json'
B4 = ROOT / 'docs/analysis/independent_4_completed_20261022'
REACH = ROOT / 'docs/analysis/reachable_rescues_20261016'
TITLE = '量子LLM支援プログラミングにおける失敗類型と介入カタログの実証研究'
NAVY, INK, MUTED = '172D44', '203348', '52677B'
TEAL, ORANGE, PURPLE, GREY = '087F8C', 'BE641C', '7553A6', '718398'
PALE, WHITE, LINE = 'F0F5F8', 'FFFFFF', 'DCE5EB'
COLORS = {'R': GREY, 'F': TEAL, 'D': ORANGE, 'B4': PURPLE}
FONT = 'IPAexGothic'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title, prs.core_properties.author = TITLE, '伊藤 悠真'
prs.core_properties.subject = 'WiNF 発表用下書き／本編約12分＋質疑用補足'
notes, records = [], []


def read_csv(path):
    return list(csv.DictReader(path.open()))


snapshot = json.loads(SNAP_PATH.read_text())
rows = read_csv(B4/'summary.csv')


def row(c, grouping='overall', group='all', cap='4'):
    return next(r for r in rows if r['condition'] == c and r['grouping'] == grouping
                and r['group'] == group and r['generation_cap'] == cap)


def added(c, grouping='overall', group='all'):
    return int(row(c, grouping, group)['successes_from_initial_failure'])


def rect(s, x, y, w, h, color=PALE):
    a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = RGBColor.from_string(color)
    a.line.fill.background()
    a._element.spPr.append(OxmlElement('a:effectLst'))
    return a


def text(s, x, y, w, h, value, size=24, bold=False, color=INK):
    a = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = a.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.025)
    tf.margin_top = tf.margin_bottom = Inches(.025)
    for i, line in enumerate(str(value).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.space_after = Pt(7); p.line_spacing = 1.12
        for r in p.runs:
            r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)
            rp = r._r.get_or_add_rPr(); rp.set('lang', 'ja-JP')
            ea = OxmlElement('a:ea'); ea.set('typeface', FONT); rp.append(ea)
    return a


def base(title, section, source='[3] 現行到達経路：10/16固定', appendix=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    rect(s, 0, 0, .16, 7.5, NAVY)
    text(s, .62, .25, 11.9, .3, f'WiNF 2026  /  {section}', 12, color=TEAL)
    text(s, .62, .82, 12.05, .9, title, 29, True, NAVY)
    rect(s, .65, 1.73, 12.0, .018, LINE)
    text(s, .65, 7.08, 11.7, .25, source + '  ｜  2026-10-28更新', 10, color=MUTED)
    text(s, 12.08, 7.03, .65, .35, f'{len(prs.slides):02}', 12, color=MUTED)
    records.append({'slide': len(prs.slides), 'title': title, 'appendix': appendix, 'source': source})
    return s


def note(s, seconds, body):
    n = len(prs.slides)
    msg = f'{n}枚目　' + (f'目安 {seconds}秒' if seconds else '質疑用・本編時間外') + '\n\n' + body
    s.notes_slide.notes_text_frame.text = msg
    notes.append(msg); records[-1]['duration_seconds'] = seconds


def callout(s, value, y=6.27, color=TEAL, size=22):
    rect(s, .65, y, 12.0, .63, PALE)
    rect(s, .65, y, .055, .63, color)
    text(s, .85, y+.08, 11.55, .48, value, size, True, color)


def card(s, x, y, w, label, value, desc, color, value_size=37):
    rect(s, x, y, w, 2.28, PALE)
    text(s, x+.2, y+.15, w-.4, .4, label, 20, True, color)
    text(s, x+.2, y+.66, w-.4, .75, value, value_size, True, color)
    text(s, x+.2, y+1.6, w-.4, .5, desc, 18)


def table(s, x, y, widths, data, row_h=.52, font=21):
    shape = s.shapes.add_table(len(data), len(widths), Inches(x), Inches(y),
                               Inches(sum(widths)), Inches(row_h*len(data)))
    t = shape.table
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i, rr in enumerate(data):
        for j, val in enumerate(rr):
            c = t.cell(i, j); c.text = str(val)
            c.margin_left = c.margin_right = Inches(.12)
            c.margin_top = c.margin_bottom = Inches(.025)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string(NAVY if i==0 else (PALE if i%2 else WHITE))
            for p in c.text_frame.paragraphs:
                p.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = FONT; r.font.size = Pt(font); r.font.bold = i==0
                    r.font.color.rgb = RGBColor.from_string(WHITE if i==0 else INK)
                    ea = OxmlElement('a:ea'); ea.set('typeface', FONT); r._r.get_or_add_rPr().append(ea)
    return shape


def bars(s, entries, maximum, suffix, y=2.3, step=.83, x=4.8, w=5.05, label_w=3.9):
    for i, (label, v, display, color) in enumerate(entries):
        yy=y+i*step
        text(s, .85, yy, label_w, .55, label, 22)
        rect(s, x, yy+.08, w, .36, PALE)
        rect(s, x, yy+.08, w*v/maximum, .36, color)
        text(s, x+w+.18, yy-.02, 2.35, .62, display, 23, True, color)
    bottom=y+(len(entries)-1)*step+.68
    for v in [0,maximum/2,maximum]:
        text(s, x+w*v/maximum-.1, bottom, .95, .35, f'{v:g}{suffix}', 12, color=MUTED)


# 13 main slides; 640 seconds (10m40s).
s=base('初期失敗コードを，どこまで救済できるか', '研究の問い', '[3][4] 検証済みの固定集計')
text(s,.8,2.1,11.8,1.4,'量子LLM支援プログラミングにおける\n失敗類型と介入カタログの実証研究',29,True)
text(s,.85,4.05,11.7,1.1,'2023TC011 伊藤 悠真　　指導教員：横山 哲郎\n南山大学 理工学部 電子情報学科',23)
callout(s,'救済率・救済対象・追加消費から，介入の候補を整理する')
note(s,20,'量子回路コードの初期失敗に対して、どの介入で何を救済できるかを調べます。修正だけでなく、独立に生成し直すB4も対照として比較します。')

s=base('実行できても，期待した量子出力とは限らない','背景と目的','[1] Self-Refine ／ [2] Self-Debugging ／ [3] 本研究の固定集計')
for x,l,v,desc,c in [(.85,'読込み','import・構文','APIや必須関数の要件',GREY),(4.95,'回路構築','回路を作れるか','構築・評価処理の実行',TEAL),(9.05,'出力検証','期待出力か','指定状態・測定分布',ORANGE)]:
    card(s,x,2.25,3.45,l,v,desc,c,value_size=28)
text(s,.9,4.95,11.7,.9,'研究の問い：どれだけ救済できるか／何を救済するか',27,True)
callout(s,'失敗類型・モデル・追加消費を併記した介入カタログへ')
note(s,45,'量子コードでは、Pythonとして読み込めても回路を作れない場合や、実行できても期待する状態や分布にならない場合があります。Self-RefineとSelf-Debuggingを背景に独自の修正条件を実装しました。先行研究の完全再現ではありません。')

s=base('同じ初期コードから，修正と再生成を比較する','比較条件','[3] B/R/F/D/P ／ [4] B4追加対照')
rect(s,.85,3.1,2.25,1.15,NAVY);text(s,1.02,3.36,1.95,.6,'初期生成 B',25,True,WHITE)
for y,c,label,desc in [(2.08,'R','自己見直し','自分のコードを見直す'),(2.94,'F','実行結果提示','エラー等の診断を提示'),(3.8,'D','自己デバッグ','診断＋説明・原因分析'),(4.66,'B4','独立再生成','前回コード・診断は渡さない')]:
    text(s,3.3,y+.1,.6,.5,'→',26,color=COLORS[c]);rect(s,4.05,y,8.3,.7,PALE)
    text(s,4.2,y+.07,3.7,.5,c+'  '+label,22,True,COLORS[c]);text(s,8.12,y+.1,4,.45,desc,18)
callout(s,'R/F/D/B4：初回込み最大4回・成功時停止／Pは別プロンプトの独立1回生成',size=19)
note(s,55,'BとR・F・Dは初期コードと初期判定を共有します。Rは外部の実行診断を与えず見直しを促し、Fは実行結果、Dは実行結果に加え説明と原因分析を促します。B4は初期コードを共有し、その失敗組に対して最大3回、課題仕様から独立再生成します。前回コードや診断を入力しません。Pは生成前注意事項を使う別の1回生成です。')

s=base('対象は1,000初期組．出典を揃えて比較する','実験設計','[3] 現行到達経路 ／ [4] B4追加対照（10/22固定）')
text(s,.85,2.05,11.7,.7,'5モデル × 10タスク × seed 0〜19 ＝ 1,000初期組',28,True)
table(s,.85,3.02,[3.2,8.25],[['項目','設定・範囲'],['モデル','Q4 / G4 / Q9 / C7 / G3（正式タグは補足）'],['タスク','Bell・GHZ・DJ 2種・BV・Grover・QFT・IQFT・QPE・Ansatz'],['生成設定','temperature 0.7 ／ 1要求の生成上限 1,024 tokens'],['集計対象','B/R/F/D/P 計5,000条件 ＋ 同じ初期組への追加対照B4']],row_h=.53,font=20)
callout(s,'保存記録＋必要な部分追試293ケースを反映．全件新生成の実験ではない',size=20)
note(s,60,'本体は5モデル、10タスク、20 seedの千組です。元の3実験は3モデルseed 0〜9の1500条件、5モデルseed 10〜19の2500条件、残り2モデルseed 0〜9の1000条件です。現在の到達経路は検証済み保存prefixと必要な追試suffixで再構成しています。B4は同じ千組への追加対照で、本体の5000条件には混ぜません。出典や時期の差は比較の限界です。')

s=base('救済率の分母は，初期失敗545組','成功と指標の定義')
card(s,.85,2.04,5.6,'初期成功','455組','無修正で保持',TEAL)
card(s,6.65,2.04,5.6,'初期失敗','545組','修正・再生成の対象',ORANGE)
text(s,.92,4.65,11.5,.7,'救済率 ＝ 初期失敗から修正で最終成功となった組数 ／ 545',24,True)
text(s,.92,5.49,11.5,.55,'主L3：読込み・構築・指定入力の出力一致 ＋ 所定の構造要件',22)
callout(s,'主L3は全入力のアルゴリズム同値性を保証しない',size=23)
note(s,45,'千組のうち455組が初期成功、545組が初期失敗です。救済率の分母はこの545組で、モデル率の単純平均ではありません。失敗類型も初期判定で固定します。B4では同じ分母を使いますが、同一コードの修正ではないため追加成功と呼びます。成功判定の主L3は指定入力と構造要件に限定されます。')

s=base('FとDは初期失敗の約3分の1を救済','主要結果')
entries=[]
for c,label in [('R','自己見直し'),('F','実行結果提示'),('D','自己デバッグ')]:
    n=added(c);entries.append((c+'  '+label,n/545*100,f'{n}組  {n/545*100:.1f}%',COLORS[c]))
bars(s,entries,50,'%',y=2.5,step=1.0)
text(s,.9,5.88,11.7,.32,'共通分母：初期失敗545組／初期成功455組の無修正保持は救済数に含めない',17,color=MUTED)
callout(s,'FとDの全体差は3組．同じ対象を救済しているか？')
note(s,60,'Rは96組で17.6%、Fは183組で33.6%、Dは180組で33.0%を救済しました。いずれも分母545組です。今回の対象ではFとDがRより多く救済しましたが、一般的な優劣は断定しません。FとDの差は3組と小さい一方、救済した対象が一致するかは別に確認する必要があります。')

s=base('救済数が近くても，救済する対象は異なる','救済対象の対応')
for x,l,n,c in [(.85,'Fのみ',43,TEAL),(4.95,'F/D共通',140,NAVY),(9.05,'Dのみ',40,ORANGE)]:card(s,x,2.1,3.45,l,f'{n}組','初期失敗545組の内訳',c)
text(s,.9,4.85,11.5,.7,'片側救済83組：両条件166履歴・398修正遷移を静的に確認',24,True)
text(s,.9,5.55,11.5,.5,'双方未救済322組／複数変更が同時に起きるため，救済原因は分離できない',20)
callout(s,'和集合223組は事後的な集合．実行前の選択方策の成功率ではない',size=21)
note(s,55,'共通救済は140組、Fのみ43組、Dのみ40組、双方未救済322組でした。片側83組の全履歴を確認しました。import、回路構築、ゲート、測定取得など複数の変更が同時に現れ、静的な差分だけで成功原因は決められません。少なくとも一方が救済した223組は、結果を見て選んだ事後的な集合です。')

s=base('初期失敗類型によって，FとDの関係が変わる','失敗類型と介入')
table(s,.85,2.12,[3.25,1.4,2.15,2.15,2.45],[['初期失敗類型','分母','R','F','D'],['import失敗',271,'24 (8.9%)','58 (21.4%)','72 (26.6%)'],['回路構築失敗',126,'32 (25.4%)','76 (60.3%)','68 (54.0%)'],['出力不一致',65,'16 (24.6%)','24 (36.9%)','18 (27.7%)']],row_h=.76,font=21)
text(s,.92,5.43,11.5,.55,'類型は介入前で固定／モデル・タスクの構成と関連する',22)
callout(s,'カタログでは全体率だけでなく，対象類型と分母を示す',size=23)
note(s,55,'import失敗ではDがFより多く救済し、回路構築失敗ではFが多いという関係でした。出力不一致でも内訳が違います。失敗類型は介入前で固定しています。ただし類型とモデルやタスクは関連するため、例えばimportなら常にDという選択規則は確定できません。率と分母を併記することが重要です。')

s=base('独立再生成B4では217組が追加成功した','追加対照B4','[3] R/F/D再構成値 ／ [4] B4完了集計')
bars(s,[(c+'  '+label,added(c)/545*100,f'{added(c)}組  {added(c)/545*100:.1f}%',COLORS[c]) for c,label in [('R','自己見直し'),('F','実行結果提示'),('D','自己デバッグ'),('B4','独立再生成')]],50,'%',y=2.05,step=.8)
text(s,.9,5.79,11.5,.35,'B4最終成功 672/1,000組（67.2%）／欠落・重複・予定外0，sanity検証成功',18)
callout(s,'同じ上限4回でも，入力情報・実要求数・出典・実行時期が異なる',size=21,color=PURPLE)
note(s,60,'B4は初期失敗から217組、39.8%が成功し、初期成功を含む最終成功は672組でした。Fより34組、Dより37組多い観測です。最大生成機会を揃えても、B4は過去コードや診断を受け取らず、seed規則や実行時期も異なります。R・F・Dは保存経路の再構成であり、診断を与えることの因果効果や一般的優劣は断定できません。')

s=base('追加消費を含めると，独立再生成も比較候補になる','追加計算コスト','[3] R/F/D再構成値 ／ [4] B4実測記録')
entries=[(c, float(row(c)['tokens_per_added_success']),f"{float(row(c)['tokens_per_added_success']):,.1f}",COLORS[c]) for c in ['R','F','D','B4']]
bars(s,entries,20000,'',y=2.1,step=.79,x=3.35,w=6.6,label_w=2.25)
text(s,.9,5.8,11.5,.35,'指標：追加tokens／追加成功．共有初期を除き，未成功例への消費も含む',19)
callout(s,'R/F/Dは再構成経路の計算値．金銭費用・実測の費用削減ではない',size=21)
note(s,50,'未成功例への消費も含む追加tokensを追加成功数で割ると、Rは18852.6、Fは9333.8、Dは10976.6、B4は4589.0でした。共有する初期生成は除いています。B4は1342回の追加要求で995815 tokensでした。入力長や出典などが異なり、R・F・Dは再構成値です。金銭費用や実測の費用削減を示すものではありません。')

s=base('3回目にも救済がある．停止回数も検討対象','修正回数と停止')
fs=read_csv(B4/'first_success.csv')
data=[['修正条件','1回目','2回目','3回目','2回打切りで失う救済']]
for c in ['R','F','D']:
    ns=[int(next(r['groups'] for r in fs if r['condition']==c and r['first_success_generation']==f'generation_{g}')) for g in [2,3,4]]
    data.append([c]+[f'{n}組' for n in ns]+[f'{ns[2]/sum(ns)*100:.1f}%'])
table(s,.85,2.25,[2.2,1.75,1.75,1.75,3.95],data,row_h=.75,font=22)
text(s,.92,5.52,11.5,.5,'失う割合の分母は「3回までに救済した組数」．初期失敗数とは異なる',20)
callout(s,'同じ観測履歴を打ち切った比較．上限別独立実験ではない',size=23)
note(s,50,'Fの修正ごとの救済は112、43、28組、Dは108、39、33組です。2回で止めると、最大3回で救済できた対象のうちFでは15.3%、Dでは18.3%を失います。この割合の分母は救済総数です。各回上限で独立に実験したものではなく、同じ履歴を打ち切った計算なので、一般的に最適な回数は確定できません。')

s=base('介入カタログには，適用範囲と限界を併記する','考察と次のアクション')
for y,n,title,desc in [(2.1,'01','何に使うか','初期失敗類型・モデル・分母・成功対象'),(3.25,'02','どれだけ使うか','未成功消費を含むtokens・要求数・停止回数'),(4.4,'03','どこまで分かるか','指定入力の主L3・出典差・共有初期とタスク依存')]:
    text(s,.85,y,1,.65,n,30,True,TEAL);text(s,2,y,3.35,.6,title,25,True);text(s,5.35,y+.04,7,.65,desc,21)
callout(s,'次：選択・停止方策を事前固定し，別seed・タスクで検証する',size=22)
note(s,45,'カタログの候補はR・F・Dに独立再生成も加えます。記載すべきなのは、対象の失敗類型やモデル、分母、成功対象、追加消費、停止条件、適用範囲です。共有初期やタスク依存があり、5000条件を独立標本とは扱いません。今後は選択と停止を事前に固定し、別seedとタスクで検証します。')

s=base('救済率に加え，「何を救済するか」を比較する','まとめ','[3][4] 検証済み固定集計')
for y,n,t,sub in [(2.04,'1','F 183組・D 180組を救済','全体率は近くても，片側救済83組の対象が異なる'),(3.36,'2','独立再生成B4は217組が追加成功','修正と再生成を，入力情報・出典差も含めて比較する'),(4.68,'3','分母・対象・追加消費をカタログに記録','選択・停止方策を固定して別条件で検証する')]:
    text(s,.85,y,.7,.7,n,34,True,TEAL);text(s,1.8,y,10.6,.6,t,27,True);text(s,1.8,y+.65,10.6,.5,sub,21)
callout(s,'今回の対象での観測であり，一般的優劣・因果効果・統計的有意差は主張しない',size=19)
note(s,40,'FとDは初期失敗の約3分の1を救済しましたが、片側だけが救済した83組がありました。B4は217組の追加成功で、修正だけではなく再生成も対照にする必要性が示されました。率、対象、追加消費を揃えて記録し、今後は事前に固定した方策を別条件で検証します。以上です。')

# Appendix slides.
s=base('補足A　モデル別の初期失敗と追加成功','モデル別内訳','[3] R/F/D ／ [4] B4',True)
model_map=[('G3','gemma3:4b'),('G4','gemma4:e4b-it-qat'),('C7','qwen2.5-coder:7b'),('Q4','qwen3.5:4b'),('Q9','qwen3.5:9b')]
data=[['略号／正式タグ','初期失敗','R','F','D','B4']]
for short,m in model_map:data.append([short+'  '+m,row('B4','model','ollama:'+m)['initial_failures']]+[added(c,'model','ollama:'+m) for c in ['R','F','D','B4']])
table(s,.7,2.12,[5.22,1.53,1.12,1.12,1.12,1.12],data,row_h=.55,font=20)
callout(s,'モデルごとの失敗分母が異なる．モデル一般の優劣は断定しない',size=22)
note(s,0,'G4ではF48組・D38組、Q4ではF78組・D90組と大小が逆転します。タグを識別子とし、表記の数値を厳密なパラメータ数と解釈しません。B4は修正ではなく独立再生成による追加成功です。')

s=base('補足B　B4とFでも成功する対象は一致しない','対応比較とP','[3] Pの既存集計 ／ [4] B4 pairwise.csv',True)
pair=next(r for r in read_csv(B4/'pairwise.csv') if r['comparator']=='F')
table(s,.85,2.13,[3,3,3,3],[['共通成功','B4のみ','Fのみ','双方未成功'],[pair['both'],pair['B4_only'],pair['F_only'],pair['neither']]],row_h=.77,font=25)
text(s,.92,4.22,11.5,.65,'B4とFの対応比較の分母：初期失敗545組',25,True)
text(s,.92,5.05,11.5,.98,'Pは独立1回生成：初期失敗から成功124組，\n初期成功に対応する失敗106組．最終成功473/1,000組．',23)
callout(s,'事後的な和集合を，実行前に使える選択方策とみなさない',size=22)
note(s,0,'B4とFは共通131、B4のみ86、Fのみ52、双方未成功276です。B4で成功するからといってFで成功する組をすべて含むわけではありません。Pは別プロンプトによる独立1回生成で、初期成功側に対応する失敗もあるためR・F・Dと同じ作用と解釈しません。')

s=base('補足C　評価範囲・経路・データ品質','検証と限界','[3][4] 保存設定・manifest・sanity report',True)
table(s,.75,2.02,[2.25,9.65],[['項目','検証範囲と限界'],['主L3','L0：コード読込み・必須関数／L1：構築・評価／L2：指定出力一致＋構造'],['出力の基準','分布：8,192 shotsの全変動距離／状態：位相整合した距離／0.05未満'],['構造・乱数','量子ビット数・測定/barrier以外の命令数／測定・transpile seed 42'],['到達経路','保存prefix＋必要な追試suffix．分岐後に到達しない旧suffixを除外'],['比較の限界','共有初期・タスク・seed依存，schema・診断文・実行時期の差'],['品質確認','B4：全1,000組，欠落・重複・予定外0．コード側NameErrorとAPI障害を区別']],row_h=.55,font=17)
callout(s,'旧120件は別設定の既報として保持し，主結果に合算しない',size=22)
note(s,0,'出力一致は全入力のアルゴリズム同値性ではありません。旧120件は2モデル、seed 0〜2、2条件、L2終了の既報です。B4の追加roundのunknown_error2件はコード側NameErrorでありAPI障害ではありません。元ログには中断時の接続エラーがありますが、保存設定で再開した全件完了の記録を基準にしています。')

s=base('補足D　参考文献と固定出典','出典','[1]–[4] 参考文献・公開用集計',True)
text(s,.85,2.02,11.6,.95,'[1] A. Madaan et al.  Self-Refine: Iterative Refinement with Self-Feedback\n      arXiv:2303.17651 (2023).',19)
text(s,.85,3.13,11.6,.95,'[2] X. Chen et al.  Teaching Large Language Models to Self-Debug\n      arXiv:2304.05128 (2023).',19)
text(s,.85,4.24,11.6,.9,'[3] 現行到達経路の固定集計（2026-10-16）\n      docs/presentation_sources/20261016T_reachable_rescues/',18)
text(s,.85,5.25,11.6,.9,'[4] B4追加実験の完了集計（2026-10-22固定）\n      docs/analysis/independent_4_completed_20261022/',18)
text(s,.85,6.38,11.6,.5,'https://github.com/yuuma-ito/2026_ito_qLLM-ver3',18,color=TEAL)
note(s,0,'出典の保存設定、source_manifest、summary、pairwise、sanityを参照してください。今回の資料作成では新規モデル要求や生成コードの実行は行っていません。')


def export_pdf():
    local = ROOT / '.cache/graduation_office/root/usr/lib'
    env = os.environ.copy(); binary = shutil.which('libreoffice')
    if not binary:
        binary = str(local/'libreoffice/program/soffice.bin')
        env['LD_LIBRARY_PATH'] = ':'.join([str(local/'libreoffice/program'), str(local/'x86_64-linux-gnu'), env.get('LD_LIBRARY_PATH','')])
        env['SAL_USE_VCLPLUGIN'] = 'svp'
    with tempfile.TemporaryDirectory(prefix='winf-slides-office-') as profile:
        command = [binary, '-env:UserInstallation='+Path(profile).as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(HERE), str(HERE/'slides.pptx')]
        pdf = HERE/'slides.pdf'
        before = pdf.stat().st_mtime_ns if pdf.exists() else None
        # A fresh LibreOffice profile may request one startup restart (exit 81).
        for attempt in range(2):
            result = subprocess.run(command, env=env, capture_output=True, text=True, timeout=90)
            if result.returncode == 81 and attempt == 0:
                continue
            result.check_returncode()
            break
        if not pdf.exists() or pdf.stat().st_mtime_ns == before:
            raise RuntimeError('PDF was not refreshed: '+result.stdout+result.stderr)
        print(result.stdout.strip())


def build():
    for path, digest in snapshot['source_sha256'].items():
        assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest() == digest, path
    assert [added(c) for c in ['R','F','D','B4']] == [96,183,180,217]
    prs.save(HERE/'slides.pptx')
    total = sum(r['duration_seconds'] for r in records)
    (HERE/'speaker_notes.txt').write_text(f'WiNF発表メモ：本編13枚・{total//60}分{total%60:02}秒（目安）、補足4枚\n発表時間は未定。ユーザー指定の10〜12分を想定。\n\n'+'\n\n'.join(notes)+'\n')
    files = [SNAP_PATH, HERE.parent/'paper.tex', B4/'summary.csv', B4/'first_success.csv', B4/'pairwise.csv', B4/'source_manifest.json', REACH/'core/repair_summary.csv']
    (HERE/'source_manifest.json').write_text(json.dumps({'built_at_utc':datetime.now(timezone.utc).isoformat(),'title':TITLE,'main_slides':13,'appendix_slides':4,'duration_seconds':total,'B4_experiment_id':snapshot['B4_experiment_id'],'slides':records,'source_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in files},'new_generations':0,'generated_code_executions':0},ensure_ascii=False,indent=2)+'\n')
    print(f'Created {len(prs.slides)} slides; main {total} seconds')


if __name__ == '__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--pdf',action='store_true');args=parser.parse_args()
    build()
    if args.pdf: export_pdf()
